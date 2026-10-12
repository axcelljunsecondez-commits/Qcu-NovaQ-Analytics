"""B8 plan report (level 1): routes, gate, report model and workbook.

Spec: docs/superpowers/specs/2026-10-09-b8-reports-pdf-excel.md §6 tests 1-4 (owner answers 2026-10-09).
The shared case goes through the public routes only (create, upload, optimize, save, select), as a user would.

Data: SYNTHETIC (rates and setups made for the test).
"""

from __future__ import annotations

import io
import math
import statistics
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.db.models import Dataset, Scenario
from backend.reports.formula_check import COL, FIRST_ROW
from backend.reports.plan_report import COVER_LABEL
from tests.excel_formula_eval import SheetEvaluator
from tests.helpers import create_user, csrf_header, login, make_sessionmaker

SHARED_SETUP = {"queue_structure": "shared_queue", "fixed_server_count": 2, "staffing_varies_by_period": False,
                "capacity_mode": "unlimited", "abandonment_mode": "not_modeled"}
CSV = b"time,lambda,mu,c\n08:00-09:00,40,15,2\n09:00-10:00,20,15,2\n"
OPTIONS = {"target_utilization": 0.85, "server_cost_per_hr": 87.0, "customer_waiting_cost": 100.0, "max_servers": 6}


def _shared_plan(client, db_engine) -> dict[str, Any]:
    create_user(db_engine, "owner@example.com", "pw")
    assert login(client, "owner@example.com", "pw") == 200
    analysis = client.post("/analyses", headers=csrf_header(client),
                           json={"name": "Counters", "queue_setup": SHARED_SETUP}).json()["analysis"]
    dataset = client.post(f"/analyses/{analysis['id']}/datasets", headers=csrf_header(client),
                          files={"file": ("rates.csv", CSV, "text/csv")}).json()["dataset"]
    segments = [{key: row[key] for key in ("time", "lambda", "mu", "c")} for row in dataset["normalized"]]
    results = client.post("/optimize/batch", headers=csrf_header(client), json={"segments": segments, **OPTIONS}).json()
    saved = client.post("/scenarios", headers=csrf_header(client), json={
        "name": "Four at eight", "analysis_id": analysis["id"], "dataset_id": dataset["id"], "results": results,
        "settings": {**OPTIONS, "calculation": {
            "schema_version": 1, "engine_version": "novaq-2026-09-system-v2", "input_segments": segments,
            "options": OPTIONS, "what_if_multiplier": 1, "calculated_at": "2026-10-13T00:00:00+00:00"}}})
    assert saved.status_code == 201, saved.text
    return {"analysis_id": analysis["id"], "dataset_id": dataset["id"], "scenario_id": saved.json()["scenario"]["id"],
            "results": results["results"], "segments": segments}


def _select(client, ids):
    response = client.post(f"/analyses/{ids['analysis_id']}/workflow/selection", headers=csrf_header(client),
                           json={"scenario_id": ids["scenario_id"]})
    assert response.status_code == 200, response.text


def _preview(client, analysis_id):
    return client.get(f"/analyses/{analysis_id}/reports/plan/preview")


def test_the_plan_report_needs_a_selected_current_plan(db_engine, client):
    ids = _shared_plan(client, db_engine)
    refused = _preview(client, ids["analysis_id"])
    assert refused.status_code == 422 and refused.json()["detail"]["code"] == "no_selected_plan"
    _select(client, ids)
    assert _preview(client, ids["analysis_id"]).status_code == 200
    # A new planning file makes the plan stale: the report says so instead of showing old numbers.
    client.post(f"/analyses/{ids['analysis_id']}/datasets", headers=csrf_header(client),
                files={"file": ("rates.csv", CSV.replace(b"40", b"41"), "text/csv")})
    stale = _preview(client, ids["analysis_id"])
    assert stale.status_code == 422 and stale.json()["detail"]["code"] == "plan_not_current"
    assert client.get(f"/analyses/{ids['analysis_id']}/reports/plan.xlsx").status_code == 422


def test_the_shared_model_reads_the_saved_plan(db_engine, client):
    ids = _shared_plan(client, db_engine)
    _select(client, ids)
    model = _preview(client, ids["analysis_id"]).json()["model"]
    results = ids["results"]
    assert model["cover"]["label"] == COVER_LABEL and model["cover"]["plan_name"] == "Four at eight"
    assert [(p["counters_now"], p["counters_plan"]) for p in model["periods"]] == [
        (row["c_current"], row["c_optimal"]) for row in results]
    assert model["periods"][0]["counters_plan"] == 4  # 40/h with 15/h each: 3 counters is 88.9 % busy, above 85 %
    assert model["key_results"]["cost_now"] is None  # 2 counters cannot keep up with 40/h
    assert model["key_results"]["cost_plan"] == pytest.approx(math.fsum(row["cost_optimal"] for row in results))
    assert model["key_results"]["counter_changes"][0] == {"time": "08:00-09:00", "from": 2, "to": 4}
    assert model["key_results"]["highest_wait_plan_minutes"] == pytest.approx(
        max(row["Wq_optimal"] for row in results) * 60)
    search = model["plan_search"]
    assert search["source"] == "stored"
    first = [row for row in search["rows"] if row["time"] == "08:00-09:00"]
    assert [row["counters"] for row in first] == [1, 2, 3, 4, 5, 6]
    assert [row["result"] for row in first][:4] == [
        "fails: counters cannot keep up, busier than the target, no finite wait",
        "fails: counters cannot keep up, busier than the target, no finite wait",
        "fails: busier than the target", "chosen"]
    assert first[3]["total_cost"] == pytest.approx(4 * 87.0 + 40 * first[3]["wait_minutes"] / 60 * 100.0)
    assert {row["staffing"] for row in model["formula_check"]} == {"Now", "Plan"}


def test_the_workbook_formulas_recompute_the_saved_numbers(db_engine, client):
    ids = _shared_plan(client, db_engine)
    _select(client, ids)
    response = client.get(f"/analyses/{ids['analysis_id']}/reports/plan.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Overview", "Your data", "Plan", "Formula check", "Plan search", "Assumptions", "Provenance"]
    sheet = workbook["Formula check"]
    evaluator = SheetEvaluator(sheet)
    checked = 0
    for r in range(FIRST_ROW, sheet.max_row + 1):
        excel = evaluator.value(f"{get_column_letter(COL['wq_x'])}{r}")
        novaq = sheet.cell(r, COL["wq_n"]).value
        if isinstance(excel, float):
            assert excel == pytest.approx(novaq, rel=1e-9, abs=1e-12)
            checked += 1
        else:
            assert excel == novaq == "Unstable"
    assert checked == 3  # 08:00 with 4 counters, 09:00 now and with the plan; 08:00 now is unstable
    text = " ".join(str(cell.value) for name in workbook.sheetnames for row in workbook[name].iter_rows() for cell in row
                    if cell.value is not None)
    assert "Decision" not in text and "recommend" not in text.lower()  # level 1 has no Decision words


def test_a_plan_saved_before_the_candidates_existed_recomputes_them(db_engine, client):
    ids = _shared_plan(client, db_engine)
    _select(client, ids)
    stored = _preview(client, ids["analysis_id"]).json()["model"]["plan_search"]["rows"]
    with make_sessionmaker(db_engine)() as db:
        scenario = db.get(Scenario, ids["scenario_id"])
        results = dict(scenario.results_json)
        results["results"] = [{k: v for k, v in row.items() if k != "candidates"} for row in results["results"]]
        scenario.results_json = results
        db.commit()
    search = _preview(client, ids["analysis_id"]).json()["model"]["plan_search"]
    assert search["source"] == "recomputed" and search["rows"] == stored


def test_a_page_from_before_the_candidates_still_saves_and_gets_them(db_engine, client):
    ids = _shared_plan(client, db_engine)
    old = [{k: v for k, v in row.items() if k != "candidates"} for row in ids["results"]]
    saved = client.post("/scenarios", headers=csrf_header(client), json={
        "name": "Old tab", "analysis_id": ids["analysis_id"], "dataset_id": ids["dataset_id"], "results": {"results": old},
        "settings": {**OPTIONS, "calculation": {
            "schema_version": 1, "engine_version": "novaq-2026-09-system-v2", "input_segments": ids["segments"],
            "options": OPTIONS, "what_if_multiplier": 1, "calculated_at": "2026-10-13T00:00:00+00:00"}}})
    assert saved.status_code == 201, saved.text
    assert all("candidates" in row for row in saved.json()["scenario"]["results"]["results"])


# ── Separate lines ──────────────────────────────────────────────────────────


def test_the_separate_model_uses_the_simulated_plan_and_checks_todays_lines(db_engine, client):
    from tests import test_b13_per_date_day as b13

    analysis_id, dataset_id, rows = b13._workspace(db_engine, "representative_day")
    with make_sessionmaker(db_engine)() as db:  # a real upload stores each line's service variance
        dataset = db.get(Dataset, dataset_id)
        dataset.normalized_json = [{**row, "variance": statistics.pvariance(row["service_samples_hours"])}
                                   for row in dataset.normalized_json]
        db.commit()
    login(client, "owner@example.com", "pw")
    out = b13._optimize(client, analysis_id)
    saved = b13._save(client, analysis_id, dataset_id, b13._settings(analysis_id, dataset_id, rows, out["engine_version"]),
                      out["schedule"])
    assert saved.status_code == 201, saved.text
    _select(client, {"analysis_id": analysis_id, "scenario_id": saved.json()["scenario"]["id"]})
    model = _preview(client, analysis_id).json()["model"]
    assert model["line_type"] == "separate"
    periods = out["schedule"]["periods"]
    assert [p["counters_plan"] for p in model["periods"]] == [p["optimal_active_lanes"] for p in periods]
    chosen = [row for row in model["plan_search"]["rows"] if row["result"] == "chosen"]
    assert len(chosen) == len(periods)
    live = [row for row in model["formula_check"] if row["live"]]
    assert live and {row["model"] for row in live} == {"Parallel M/G/1"}
    workbook = load_workbook(io.BytesIO(client.get(f"/analyses/{analysis_id}/reports/plan.xlsx").content))
    sheet = workbook["Formula check"]
    evaluator = SheetEvaluator(sheet)
    for r in range(FIRST_ROW, FIRST_ROW + len(live)):
        excel = evaluator.value(f"{get_column_letter(COL['wq_x'])}{r}")
        assert excel == pytest.approx(sheet.cell(r, COL["wq_n"]).value, rel=1e-9, abs=1e-12)


def test_another_user_cannot_read_the_report(db_engine, client):
    ids = _shared_plan(client, db_engine)
    _select(client, ids)
    client.post("/auth/logout", headers=csrf_header(client))
    create_user(db_engine, "other@example.com", "pw")
    login(client, "other@example.com", "pw")
    assert _preview(client, ids["analysis_id"]).status_code == 404
    assert client.get(f"/analyses/{ids['analysis_id']}/reports/plan.xlsx").status_code == 404


def test_recomputed_counter_counts_are_used_only_when_they_choose_the_saved_plan():
    from backend.reports.plan_report import _shared_candidates

    segment = {"time": "08:00-09:00", "lambda": 40.0, "mu": 15.0, "c": 2}
    rows, source = _shared_candidates({"c_optimal": 4}, segment, OPTIONS)
    assert source == "recomputed" and [row["servers"] for row in rows or []] == [1, 2, 3, 4, 5, 6]
    assert _shared_candidates({"c_optimal": 5}, segment, OPTIONS) == (None, "unavailable")
