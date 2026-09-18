"""Three-sheet upload: preview, confirm-before-replace, apply, export (API).

Spec: docs/superpowers/specs/2026-09-18-upload-template-setup-derivation.md.
"""
from __future__ import annotations

import io

import openpyxl

from backend.db.models import AnalysisProject, Dataset
from tests.helpers import clear_cookies, create_user, csrf_header, login, make_sessionmaker
from tests.test_setup_derivation import (
    STAFF_COLUMNS,
    expected_novamart_setup,
    novamart_events,
    novamart_workbook,
    workbook_bytes,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SAVED_SEPARATE = {
    "queue_structure": "separate_queues", "fixed_server_count": 2, "staffing_varies_by_period": False,
    "capacity_mode": "unlimited", "abandonment_mode": "not_modeled", "queue_ids": ["lane_a", "lane_b"],
    "segments": [],
}


def _login(db_engine, client, email="up@example.com"):
    create_user(db_engine, email, "pw")
    login(client, email, "pw")


def _analysis(client, setup=None) -> dict:
    body = {"name": "NovaMart"} if setup is None else {"name": "NovaMart", "queue_setup": setup}
    response = client.post("/analyses", headers=csrf_header(client), json=body)
    assert response.status_code == 201, response.text
    return response.json()["analysis"]


def _post(client, path, data, filename="novamart.xlsx", params=None):
    media = XLSX if filename.endswith(".xlsx") else "text/csv"
    return client.post(path, headers=csrf_header(client), params=params or {},
                       files={"file": (filename, data, media)})


def _preview(client, analysis_id, data, filename="novamart.xlsx"):
    return _post(client, f"/analyses/{analysis_id}/datasets/preview", data, filename)


def _upload(client, analysis_id, data, filename="novamart.xlsx", apply=None):
    params = {} if apply is None else {"apply_setup": str(apply).lower()}
    return _post(client, f"/analyses/{analysis_id}/datasets", data, filename, params)


def _stored(db_engine, analysis_id):
    with make_sessionmaker(db_engine)() as db:
        analysis = db.get(AnalysisProject, analysis_id)
        count = db.query(Dataset).filter(Dataset.analysis_id == analysis_id).count()
        return analysis.queue_setup_json, count


def test_preview_derives_setup_and_stores_nothing(db_engine, client):
    _login(db_engine, client)
    analysis = _analysis(client)
    before = _stored(db_engine, analysis["id"])
    response = _preview(client, analysis["id"], novamart_workbook())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["mode"] == "multi_sheet"
    assert body["errors"] == []
    assert body["derived_setup"] == expected_novamart_setup()
    assert body["saved_setup"]["queue_structure"] == "unknown"
    assert body["needs_confirmation"] is False
    assert {item["field"] for item in body["diff"]} >= {"queue_structure", "queue_ids", "segments", "breaks"}
    assert len(body["staff"]) == 5 and len(body["breaks"]) == 15
    assert _stored(db_engine, analysis["id"]) == before


def test_apply_setup_on_new_analysis_saves_setup_and_dataset_together(db_engine, client):
    _login(db_engine, client)
    analysis = _analysis(client)
    response = _upload(client, analysis["id"], novamart_workbook(), apply=True)
    assert response.status_code == 201, response.text
    dataset = response.json()["dataset"]
    assert dataset["validation"]["period_basis"] == "representative_day"
    assert {row["time"] for row in dataset["normalized"]} == {s["id"] for s in expected_novamart_setup()["segments"]}
    setup, count = _stored(db_engine, analysis["id"])
    assert setup == expected_novamart_setup() and count == 1
    assert client.get(f"/analyses/{analysis['id']}").json()["analysis"]["setup_status"] == "ready_for_aggregate"


def test_saved_different_setup_changes_only_after_confirmation(db_engine, client):
    _login(db_engine, client)
    analysis = _analysis(client, SAVED_SEPARATE)
    saved = _stored(db_engine, analysis["id"])[0]
    preview = _preview(client, analysis["id"], novamart_workbook()).json()
    assert preview["needs_confirmation"] is True
    assert _stored(db_engine, analysis["id"]) == (saved, 0)
    # Without apply_setup the upload behaves as today: first sheet, saved Setup.
    legacy = _upload(client, analysis["id"], novamart_workbook())
    assert legacy.status_code == 422
    assert legacy.json()["detail"] == (
        "Uploaded queue IDs are not configured for this Analysis: "
        "cashier_1, cashier_2, cashier_3, cashier_4, cashier_5.")
    assert _stored(db_engine, analysis["id"]) == (saved, 0)
    confirmed = _upload(client, analysis["id"], novamart_workbook(), apply=True)
    assert confirmed.status_code == 201, confirmed.text
    # Staffing varies, so the saved fixed_server_count is kept (spec: Derived Setup).
    assert _stored(db_engine, analysis["id"]) == ({**expected_novamart_setup(), "fixed_server_count": 2}, 1)


def test_apply_setup_rejects_legacy_files_and_derivation_errors(db_engine, client):
    _login(db_engine, client)
    analysis = _analysis(client)
    legacy = workbook_bytes({"Data Entry": novamart_events()})
    response = _upload(client, analysis["id"], legacy, apply=True)
    assert response.status_code == 422
    assert response.json()["detail"] == "apply_setup requires a workbook with an events sheet."
    bad = novamart_workbook(staff=[("cashier_1", "05:00", "17:00")])
    response = _upload(client, analysis["id"], bad, apply=True)
    assert response.status_code == 422
    assert response.json()["detail"] == "The staff sheet has no row for: cashier_2, cashier_3, cashier_4, cashier_5."
    assert _stored(db_engine, analysis["id"])[1] == 0
    assert _stored(db_engine, analysis["id"])[0]["queue_structure"] == "unknown"


def test_legacy_workbook_and_csv_preview_as_legacy_and_upload_as_before(db_engine, client):
    _login(db_engine, client)
    analysis = _analysis(client, {**SAVED_SEPARATE, "queue_ids": ["cashier_1", "cashier_2", "cashier_3",
                                                                  "cashier_4", "cashier_5"], "fixed_server_count": 5})
    legacy = workbook_bytes({"Data Entry": novamart_events()})
    assert _preview(client, analysis["id"], legacy).json() == {"mode": "legacy"}
    csv = novamart_events().to_csv(index=False).encode()
    assert _preview(client, analysis["id"], csv, "events.csv").json() == {"mode": "legacy"}
    saved = _stored(db_engine, analysis["id"])[0]
    for data, name in ((legacy, "novamart.xlsx"), (csv, "events.csv")):
        response = _upload(client, analysis["id"], data, name)
        assert response.status_code == 201, response.text
    assert _stored(db_engine, analysis["id"]) == (saved, 2)


def test_setup_workbook_export(db_engine, client):
    _login(db_engine, client)
    analysis = _analysis(client)
    assert _upload(client, analysis["id"], novamart_workbook(), apply=True).status_code == 201
    response = client.get(f"/analyses/{analysis['id']}/setup/workbook")
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["staff", "breaks"]
    assert next(book["staff"].iter_rows(values_only=True)) == tuple(STAFF_COLUMNS)
    assert book["breaks"].max_row == 16
    shared = _analysis(client, {"queue_structure": "shared_queue", "fixed_server_count": 2})
    response = client.get(f"/analyses/{shared['id']}/setup/workbook")
    assert response.status_code == 422
    assert response.json()["detail"] == "The setup workbook is available for separate queues."
    clear_cookies(client)
    _login(db_engine, client, "other@example.com")
    assert client.get(f"/analyses/{analysis['id']}/setup/workbook").status_code == 404
    assert _preview(client, analysis["id"], novamart_workbook()).status_code == 404


def test_break_name_is_optional_and_bounded(db_engine, client):
    _login(db_engine, client)
    setup = {**SAVED_SEPARATE, "breaks": [
        {"queue_id": "lane_a", "scheduled_start_time": "09:00:00", "duration_minutes": 15},
        {"queue_id": "lane_b", "scheduled_start_time": "10:00:00", "duration_minutes": 15, "break_name": "Snack"}]}
    analysis = _analysis(client, setup)
    breaks = client.get(f"/analyses/{analysis['id']}").json()["analysis"]["queue_setup"]["breaks"]
    assert [entry.get("break_name") for entry in breaks] == [None, "Snack"]
    too_long = {**setup, "breaks": [{**setup["breaks"][0], "break_name": "x" * 51}]}
    response = client.patch(f"/analyses/{analysis['id']}", headers=csrf_header(client), json={"queue_setup": too_long})
    assert response.status_code == 422

