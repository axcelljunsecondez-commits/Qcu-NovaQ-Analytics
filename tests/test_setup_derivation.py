"""Three-sheet upload (events, staff, breaks) -> derived Setup, diff, and export.

Spec: docs/superpowers/specs/2026-09-18-upload-template-setup-derivation.md.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from backend.api.analysis_schemas import QueueSetup
from backend.data import uploads
from backend.data.setup_derivation import SetupExportError, derive_setup, setup_diff, setup_workbook

FIXTURE = Path(__file__).parent / "fixtures" / "novamart_events.csv"
CASHIERS = [f"cashier_{n}" for n in range(1, 6)]
EVENT_COLUMNS = ["arrival_time", "service_start", "service_end", "queue_id"]
STAFF_COLUMNS = ["queue_id", "shift_start", "shift_end"]
BREAK_COLUMNS = ["queue_id", "break_name", "start", "minutes"]
NOVAMART_STAFF = [(q, "05:00", "17:00") for q in CASHIERS[:3]] + [(q, "06:00", "18:00") for q in CASHIERS[3:]]
# Order of the draft workbook's ``setup_to_enter`` break table.
NOVAMART_BREAKS = (
    [(q, "", "07:00", 30) for q in CASHIERS[:3]] + [(q, "", "07:30", 30) for q in CASHIERS[3:]]
    + [(q, "", "11:00", 60) for q in CASHIERS[:3]] + [(q, "", "12:00", 60) for q in CASHIERS[3:]]
    + [(q, "", "15:00", 15) for q in CASHIERS[:3]] + [(q, "", "15:15", 15) for q in CASHIERS[3:]]
)


def novamart_events() -> pd.DataFrame:
    """The 1,600 NovaMart events (placeholder dates 2000-01-03 ... 2000-01-16)."""
    raw = pd.read_csv(FIXTURE)
    base = pd.Timestamp("2000-01-03") + pd.to_timedelta(raw["day"], unit="D")
    return pd.DataFrame({column: base + pd.to_timedelta(raw[source], unit="min")
                         for column, source in (("arrival_time", "arrival"), ("service_start", "start"),
                                                ("service_end", "end"))}).assign(queue_id=raw["queue"])


def frame(columns, rows) -> pd.DataFrame:
    return pd.DataFrame([list(row) for row in rows], columns=columns)


def sheets(events=None, staff=NOVAMART_STAFF, breaks=NOVAMART_BREAKS) -> dict:
    return {
        "events": novamart_events() if events is None else events,
        "staff": None if staff is None else frame(STAFF_COLUMNS, staff),
        "breaks": None if breaks is None else frame(BREAK_COLUMNS, breaks),
    }


def workbook_bytes(named: dict[str, pd.DataFrame | tuple[list, list]]) -> bytes:
    """Write sheets in order; values are frames or (columns, rows)."""
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name, content in named.items():
        sheet = book.create_sheet(name)
        columns, rows = ((list(content.columns), content.itertuples(index=False))
                         if isinstance(content, pd.DataFrame) else content)
        sheet.append(list(columns))
        for row in rows:
            sheet.append([value.to_pydatetime() if isinstance(value, pd.Timestamp) else value for value in row])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def novamart_workbook(staff=NOVAMART_STAFF, breaks=NOVAMART_BREAKS) -> bytes:
    named: dict = {"events": novamart_events()}
    if staff is not None:
        named["staff"] = (STAFF_COLUMNS, staff)
    if breaks is not None:
        named["breaks"] = (BREAK_COLUMNS, breaks)
    return workbook_bytes(named)


def expected_novamart_setup() -> dict:
    """The draft ``setup_to_enter`` sheet, restated."""
    segments = []
    for hour in range(5, 18):
        active = CASHIERS[:3] if hour == 5 else CASHIERS[3:] if hour == 17 else CASHIERS
        segments.append({"id": f"{hour:02d}:00-{hour + 1:02d}:00", "start_time": f"{hour:02d}:00:00",
                         "end_time": f"{hour + 1:02d}:00:00", "active_queue_ids": list(active)})
    return QueueSetup.model_validate({
        "queue_structure": "separate_queues",
        "fixed_server_count": None,
        "staffing_varies_by_period": True,
        "capacity_mode": "unlimited",
        "abandonment_mode": "not_modeled",
        "segments": segments,
        "queue_ids": CASHIERS,
        "breaks": [{"queue_id": q, "scheduled_start_time": f"{start}:00", "duration_minutes": minutes,
                    "break_name": None} for q, _, start, minutes in NOVAMART_BREAKS],
        "event_period_basis": "representative_day",
    }).model_dump(mode="json")


def small_events(rows, dates=("2026-08-03", "2026-08-04")) -> pd.DataFrame:
    """rows: (queue, arrival HH:MM, service minutes); repeated on every date."""
    out = []
    for date in dates:
        for queue, clock, minutes in rows:
            start = pd.Timestamp(f"{date} {clock}")
            out.append((start, start, start + pd.Timedelta(minutes=minutes), queue))
    return frame(EVENT_COLUMNS, out)


SMALL = [("a", "08:10", 5), ("b", "08:40", 5), ("a", "09:20", 5), ("b", "11:10", 5)]
SMALL_STAFF = [("a", "08:00", "12:00"), ("b", "08:00", "12:00")]


def small(staff=SMALL_STAFF, breaks=(), events=None) -> dict:
    return sheets(events=small_events(SMALL) if events is None else events, staff=staff,
                  breaks=list(breaks) or None)


# --- derivation -----------------------------------------------------------------------------


def test_novamart_three_sheet_file_derives_the_setup_to_enter_setup():
    result = derive_setup(sheets(), None)
    assert result.errors == []
    assert result.setup == expected_novamart_setup()
    assert novamart_events()["queue_id"].iloc[0] == "cashier_2"      # staff order wins over events order
    assert len(result.staff_rows) == 5 and len(result.break_rows) == 15
    assert [row["label"] for row in result.break_rows if row["queue_id"] == "cashier_1"] == [
        "Break 1", "Break 2", "Break 3"]
    assert result.records                                            # dry-run normalization succeeded


def test_without_staff_uses_first_appearance_fixed_staffing_and_last_arrival_hour():
    result = derive_setup(sheets(staff=None, breaks=None), None)
    assert result.errors == []
    setup = result.setup
    assert setup["queue_ids"] == ["cashier_2", "cashier_1", "cashier_3", "cashier_4", "cashier_5"]
    assert setup["staffing_varies_by_period"] is False and setup["fixed_server_count"] == 5
    assert [s["id"] for s in setup["segments"]][0] == "05:00-06:00"
    assert setup["segments"][-1]["id"] == "17:00-18:00"                # last arrival 17:59 -> 18:00
    assert len(setup["segments"]) == 13
    assert all(s["active_queue_ids"] is None for s in setup["segments"])
    assert setup["breaks"] == []


def test_headers_only_staff_and_breaks_sheets_count_as_absent():
    data = workbook_bytes({"Events": novamart_events(), "staff": (STAFF_COLUMNS, []),
                           " Breaks ": (BREAK_COLUMNS, [])})
    parsed = uploads.parse_upload_workbook("novamart.xlsx", data)
    assert parsed is not None and parsed["staff"] is None and parsed["breaks"] is None
    assert derive_setup(parsed, None).setup == derive_setup(sheets(staff=None, breaks=None), None).setup


def test_non_hour_shift_boundaries_add_segments():
    events = small_events([("a", "08:10", 5), ("a", "11:20", 5), ("b", "08:40", 5), ("b", "11:50", 5)])
    result = derive_setup(small(staff=[("a", "08:00", "11:30"), ("b", "08:30", "12:00")], events=events), None)
    assert result.errors == []
    assert [(s["id"], s["active_queue_ids"]) for s in result.setup["segments"]] == [
        ("08:00-08:30", ["a"]), ("08:30-09:00", ["a", "b"]), ("09:00-10:00", ["a", "b"]),
        ("10:00-11:00", ["a", "b"]), ("11:00-11:30", ["a", "b"]), ("11:30-12:00", ["b"])]
    assert result.setup["staffing_varies_by_period"] is True


def test_single_date_uses_per_date_and_fixed_staff_keeps_active_lists():
    result = derive_setup(small(events=small_events(SMALL, dates=("2026-08-03",))), None)
    assert result.errors == []
    assert result.setup["event_period_basis"] == "per_date"
    assert result.setup["staffing_varies_by_period"] is False
    assert result.setup["fixed_server_count"] == 2
    assert all(s["active_queue_ids"] == ["a", "b"] for s in result.setup["segments"])


def test_saved_capacity_and_abandonment_are_kept_unknown_defaults_apply():
    fresh = derive_setup(small(), {"queue_structure": "unknown"}).setup
    assert (fresh["capacity_mode"], fresh["abandonment_mode"]) == ("unlimited", "not_modeled")
    saved = {"queue_structure": "shared_queue", "fixed_server_count": 2, "capacity_mode": "finite",
             "total_system_capacity": 9, "abandonment_mode": "modeled", "patience_rate_per_hour": 2.0}
    kept = derive_setup(small(), saved)
    assert (kept.setup["capacity_mode"], kept.setup["total_system_capacity"]) == ("finite", 9)
    assert (kept.setup["abandonment_mode"], kept.setup["patience_rate_per_hour"]) == ("modeled", 2.0)
    # The existing ingestion rule is returned verbatim by the dry run.
    assert kept.errors == ["Separate-queue analysis currently requires unlimited capacity."]


# --- validation: exact messages, nothing dropped ----------------------------------------------


@pytest.mark.parametrize(("change", "message"), [
    ({"events": frame(EVENT_COLUMNS[:3], [])}, "The events sheet is missing columns: queue_id."),
    ({"staff": frame(["queue_id", "shift_start"], [("a", "08:00")])},
     "The staff sheet is missing columns: shift_end."),
    ({"breaks": frame(["queue_id", "start"], [("a", "09:00")])}, "The breaks sheet is missing columns: minutes."),
    ({"events": frame(EVENT_COLUMNS, [])}, "The events sheet has no rows."),
])
def test_workbook_structure_messages(change, message):
    assert derive_setup({**small(), **change}, None).errors == [message]


@pytest.mark.parametrize(("staff", "message"), [
    ([("", "08:00", "12:00"), ("a", "08:00", "12:00"), ("b", "08:00", "12:00")], "Staff sheet row 2: queue_id is empty."),
    ([("a", "8am", "12:00"), ("b", "08:00", "12:00")], "Staff sheet row 2: shift_start must be a time in HH:MM."),
    ([("a", "08:00", ""), ("b", "08:00", "12:00")], "Staff sheet row 2: shift_end must be a time in HH:MM."),
    ([("a", "12:00", "08:00"), ("b", "08:00", "12:00")], "Staff sheet row 2: shift_start must be before shift_end."),
    ([("a", "08:00", "12:00"), ("b", "08:00", "12:00"), ("a", "08:00", "12:00")],
     "The staff sheet lists a more than once (rows 2 and 4)."),
    ([("a", "08:00", "12:00")], "The staff sheet has no row for: b."),
    ([("a", "08:00", "12:00"), ("b", "08:00", "12:00"), ("c", "08:00", "12:00")],
     "The staff sheet lists queues with no events: c."),
])
def test_staff_messages(staff, message):
    assert derive_setup(small(staff=staff), None).errors == [message]


@pytest.mark.parametrize(("breaks", "message"), [
    ([("z", "", "09:00", 15)], "Breaks sheet row 2: queue_id z is not in the events sheet."),
    ([("a", "", "nine", 15)], "Breaks sheet row 2: start must be a time in HH:MM."),
    ([("a", "", "09:00", 0)], "Breaks sheet row 2: minutes must be a positive whole number."),
    ([("a", "", "09:00", 1.5)], "Breaks sheet row 2: minutes must be a positive whole number."),
    ([("a", "", "09:00", "x")], "Breaks sheet row 2: minutes must be a positive whole number."),
    ([("a", "L" * 51, "09:00", 15)], "Breaks sheet row 2: break_name must be at most 50 characters."),
    ([("a", "", "11:45", 30)], "Breaks sheet row 2: the a break at 11:45 for 30 minutes is outside its shift 08:00–12:00."),
    ([("a", "", "09:00", 30), ("b", "", "09:00", 30), ("a", "", "09:15", 15)],
     "Breaks sheet rows 2 and 4 overlap for a."),
])
def test_break_messages(breaks, message):
    assert derive_setup(small(breaks=breaks), None).errors == [message]


def test_break_outside_operating_day_without_staff():
    result = derive_setup(small(staff=None, breaks=[("a", "", "11:45", 30)]), None)
    assert result.errors == [
        "Breaks sheet row 2: the a break at 11:45 for 30 minutes is outside the operating day 08:00–12:00."]


def test_only_arrivals_outside_the_shift_are_errors():
    ends_late = small_events([("a", "09:50", 30), ("b", "08:40", 5)])      # service ends 10:20, after the shift
    ok = derive_setup(small(staff=[("a", "08:00", "10:00"), ("b", "08:00", "12:00")], events=ends_late), None)
    assert ok.errors == []
    late = small_events([("a", "09:50", 5), ("a", "10:05", 5), ("b", "08:40", 5)])
    bad = derive_setup(small(staff=[("a", "08:00", "10:00"), ("b", "08:00", "12:00")], events=late), None)
    assert bad.errors == ["2 events for a arrive outside its shift 08:00–10:00."]


def test_events_at_23_without_staff_and_segment_without_queue():
    late = derive_setup(small(staff=None, events=small_events([("a", "08:10", 5), ("b", "23:05", 5)])), None)
    assert late.errors == ["Events that arrive at or after 23:00 cannot be placed in an hourly segment."]
    events = small_events([("a", "08:10", 5), ("b", "11:10", 5)])
    gap = derive_setup(small(staff=[("a", "08:00", "10:00"), ("b", "11:00", "12:00")], events=events), None)
    assert gap.errors == ["Segment 10:00-11:00 has no queue on shift."]


def test_unparsable_arrival_uses_existing_ingestion_message():
    events = small_events(SMALL).astype({"arrival_time": object})
    events.loc[1, "arrival_time"] = "not a time"
    assert derive_setup(small(events=events), None).errors == ["Row 2: arrival_time must be a valid timestamp."]


def test_invalid_derived_setup_reports_queue_setup_reason():
    saved = {"queue_structure": "shared_queue", "fixed_server_count": 1, "capacity_mode": "finite",
             "total_system_capacity": 1, "abandonment_mode": "not_modeled"}
    result = derive_setup(small(), saved)
    assert result.setup is None
    assert result.errors == [
        "The derived Setup is invalid: Total system capacity must include and be at least the server count."]


def test_errors_are_collected_together_and_rows_are_not_dropped():
    result = derive_setup(small(staff=[("a", "12:00", "08:00"), ("b", "08:00", "12:00")],
                                breaks=[("b", "", "nine", 15), ("b", "", "09:00", 15)]), None)
    assert result.errors == ["Staff sheet row 2: shift_start must be before shift_end.",
                             "Breaks sheet row 2: start must be a time in HH:MM."]
    assert result.setup is None
    assert len(result.staff_rows) == 2 and len(result.break_rows) == 2


# --- diff and export --------------------------------------------------------------------------


def test_setup_diff_lists_only_changed_fields():
    derived = derive_setup(sheets(), None).setup
    saved = {**derived, "event_period_basis": "per_date", "breaks": []}
    assert [item["field"] for item in setup_diff(saved, derived)] == ["event_period_basis", "breaks"]
    assert setup_diff(derived, derived) == []


def _reupload(exported: bytes, events: pd.DataFrame) -> dict:
    book = openpyxl.load_workbook(io.BytesIO(exported))
    named: dict = {"events": events}
    for name in book.sheetnames:
        rows = list(book[name].iter_rows(values_only=True))
        named[name] = (list(rows[0]), rows[1:])
    parsed = uploads.parse_upload_workbook("again.xlsx", workbook_bytes(named))
    assert parsed is not None
    setup = derive_setup(parsed, None).setup
    assert setup is not None
    return setup


def test_export_round_trip_with_staff():
    derived = derive_setup(sheets(), None).setup
    exported = setup_workbook(derived)
    assert openpyxl.load_workbook(io.BytesIO(exported)).sheetnames == ["staff", "breaks"]
    assert _reupload(exported, novamart_events()) == derived


def test_export_round_trip_without_staff_writes_headers_only():
    derived = derive_setup(sheets(staff=None, breaks=[("cashier_1", "Lunch", "11:00", 60)]), None).setup
    exported = setup_workbook(derived)
    staff_rows = list(openpyxl.load_workbook(io.BytesIO(exported))["staff"].iter_rows(values_only=True))
    assert staff_rows == [tuple(STAFF_COLUMNS)]
    again = _reupload(exported, novamart_events())
    assert again == derived and again["breaks"][0]["break_name"] == "Lunch"


def test_export_errors_are_exact():
    shared = QueueSetup.model_validate({"queue_structure": "shared_queue", "fixed_server_count": 2}).model_dump(mode="json")
    with pytest.raises(SetupExportError, match="^The setup workbook is available for separate queues.$"):
        setup_workbook(shared)
    split = QueueSetup.model_validate({
        "queue_structure": "separate_queues", "staffing_varies_by_period": True, "queue_ids": ["a", "b"],
        "segments": [{"id": f"s{n}", "start_time": f"{7 + n:02d}:00:00", "end_time": f"{8 + n:02d}:00:00",
                      "active_queue_ids": ["a", "b"] if n != 2 else ["b"]} for n in (1, 2, 3)],
    }).model_dump(mode="json")
    with pytest.raises(SetupExportError, match="^Cannot export a: it has more than one shift.$"):
        setup_workbook(split)


# --- reading: every sheet keeps the upload limits -------------------------------------------------


def test_multi_sheet_mode_detection_and_legacy():
    assert uploads.parse_upload_workbook("x.csv", b"a,b\n1,2\n") is None
    legacy = workbook_bytes({"Data Entry": (EVENT_COLUMNS, [])})
    assert uploads.parse_upload_workbook("x.xlsx", legacy) is None
    with pytest.raises(uploads.UploadError, match="^The workbook has more than one sheet named events.$"):
        uploads.parse_upload_workbook("x.xlsx", workbook_bytes({"events": (EVENT_COLUMNS, []),
                                                                 " EVENTS": (EVENT_COLUMNS, [])}))


def test_upload_limits_apply_to_every_sheet():
    events = small_events(SMALL)
    too_many_rows = workbook_bytes({"events": events.head(2), "staff": (STAFF_COLUMNS, [("a", "08:00", "12:00")] * 5)})
    with pytest.raises(uploads.UploadError, match="too many rows"):
        uploads.parse_upload_workbook("x.xlsx", too_many_rows, max_rows=3)
    wide = workbook_bytes({"events": events.head(2), "breaks": (BREAK_COLUMNS + ["x1", "x2"], [])})
    with pytest.raises(uploads.UploadError, match="too many columns"):
        uploads.parse_upload_workbook("x.xlsx", wide, max_columns=5)
    long_text = workbook_bytes({"events": events.head(2), "breaks": (BREAK_COLUMNS, [("a", "L" * 40, "09:00", 15)])})
    with pytest.raises(uploads.UploadError, match="text that is too long"):
        uploads.parse_upload_workbook("x.xlsx", long_text, max_cell_chars=30)
    sheets3 = workbook_bytes({"events": events.head(2), "staff": (STAFF_COLUMNS, []), "breaks": (BREAK_COLUMNS, [])})
    with pytest.raises(uploads.UploadError, match="too many worksheets"):
        uploads.parse_upload_workbook("x.xlsx", sheets3, xlsx_max_worksheets=2)
