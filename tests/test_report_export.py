"""Unit tests for backend.reports.report_export."""

from __future__ import annotations

import io

import openpyxl
import pandas as pd

from backend.reports.report_export import (
    _exec_summary_bullets,
    generate_excel_report,
    generate_pdf_report,
)


def _comparison_df() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "time": "08:00-09:00",
                "c_current": 2,
                "rho_current": 0.85,
                "Wq_current": 0.15,
                "c_optimal": 3,
                "rho_optimal": 0.65,
                "Wq_optimal": 0.08,
                "delta_c": 1,
            },
            {
                "time": "09:00-10:00",
                "c_current": 2,
                "rho_current": 0.92,
                "Wq_current": 0.22,
                "c_optimal": 1,
                "rho_optimal": 0.55,
                "Wq_optimal": 0.10,
                "delta_c": -1,
            },
        ]
    )


def _kpis() -> dict:
    return {
        "avg_waiting_time": 0.15,
        "avg_waiting_optimized": 0.08,
        "avg_utilization": 0.65,
        "avg_utilization_optimized": 0.43,
        "total_savings": 12000,
    }


def test_pdf_report_api_positional_call_without_segment_df() -> None:
    """The active API route passes arguments positionally and never segment_df."""
    buf = generate_pdf_report(_kpis(), _kpis(), _comparison_df(), ["Add a server."])
    assert isinstance(buf, io.BytesIO)
    assert buf.getvalue().startswith(b"%PDF")


def test_excel_report_api_positional_call_without_segment_df() -> None:
    """The active API route passes arguments positionally and never segment_df."""
    buf = generate_excel_report(_comparison_df(), _kpis())
    assert isinstance(buf, io.BytesIO)
    assert buf.getvalue().startswith(b"PK\x03\x04")


def test_excel_report_includes_staffing_summary_and_status_columns() -> None:
    buf = generate_excel_report(_comparison_df(), _kpis())
    wb = openpyxl.load_workbook(buf)

    summary_labels = [wb["Summary"].cell(row=r, column=1).value for r in range(1, 25)]
    summary_values = [wb["Summary"].cell(row=r, column=2).value for r in range(1, 25)]
    assert "Staffing Adjustment Summary" in summary_labels
    assert "Cashier-Hour Formula" in summary_labels
    assert "1 removed - 1 added" in summary_values
    assert "Schedule Action" in summary_labels
    assert "Add 1 cashier: 08:00-09:00" in summary_values
    assert "Reduce 1 cashier: 09:00-10:00" in summary_values
    assert "Utilization Status Legend" in summary_labels

    ws = wb["Segments"]
    headers = [cell.value for cell in ws[1]]
    assert "rho_current_status" in headers
    assert "rho_optimal_status" in headers
    assert ws.cell(row=2, column=headers.index("rho_current_status") + 1).value == "Peak"
    assert ws.cell(row=3, column=headers.index("rho_current_status") + 1).value == "Critical"
    assert ws.cell(row=2, column=headers.index("rho_optimal_status") + 1).value == "Normal"
    assert ws.cell(row=3, column=headers.index("rho_optimal_status") + 1).value == "Lean"


def test_exec_summary_bullets_full_pair() -> None:
    bullets = _exec_summary_bullets(_kpis(), _kpis())
    assert bullets == [
        "• Average customer wait: 9.0 min → 4.8 min (optimized)",
        "• Estimated daily savings: ₱12,000",
        "• Utilization improvement: 65% → 43%",
    ]


def test_exec_summary_bullets_current_only() -> None:
    bullets = _exec_summary_bullets(
        {"avg_waiting_time": 0.15, "avg_utilization": 0.65}, {}
    )
    assert bullets == [
        "• Average customer wait: 9.0 min (current)",
        "• Utilization improvement: 65% (current)",
    ]


def test_exec_summary_bullets_none() -> None:
    bullets = _exec_summary_bullets({}, {})
    assert bullets == [
        "• Average customer wait: N/A",
        "• Utilization improvement: N/A",
    ]
