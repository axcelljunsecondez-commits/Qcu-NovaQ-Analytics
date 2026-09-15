"""Upload templates, examples, field guide, and header aliases.

Single source of truth for what NovaQ accepts: every template column is drawn
from the real ingestion contract (``backend/data/ingestion.py`` and
``backend/data/analysis_ingestion.py``). Rule: parser requirements =
template requirements = validation requirements.
"""

from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from backend.data.analysis_ingestion import EVENT_REQUIRED
from backend.data.ingestion import HEADER_ALIASES, OPTIONAL_COLUMNS, REQUIRED_COLUMNS, normalize_header_aliases

STRUCTURES = ("shared_queue", "separate_queues")
SCHEMAS = ("aggregate", "events")

# Canonical column order per template. Separate aggregate appends the Queue ID
# and variance columns the separate path requires; event schemas use the exact
# customer-event contract.
TEMPLATE_COLUMNS: dict[tuple[str, str], list[str]] = {
    ("shared_queue", "aggregate"): list(REQUIRED_COLUMNS),
    ("separate_queues", "aggregate"): ["time", "queue_id", "lambda", "mu", "c", "variance"],
    ("shared_queue", "events"): ["arrival_time", "service_start", "service_end"],
    ("separate_queues", "events"): ["arrival_time", "service_start", "service_end", "queue_id"],
}

FIELD_GUIDE: list[dict[str, str]] = [
    {
        "field": "time",
        "meaning": "Time period label for one row of aggregate counts and rates.",
        "required": "Required for aggregate uploads",
        "example": "08:00-09:00",
        "notes": "Must not be empty. NovaQ groups these labels into operating segments.",
    },
    {
        "field": "queue_id",
        "meaning": "Which waiting line the row belongs to. queue_1 means Cashier 1's line.",
        "required": "Required for separate queues; not used for a shared queue",
        "example": "queue_1",
        "notes": "Every value must match a service line configured in Setup.",
    },
    {
        "field": "lambda",
        "meaning": "Average customer arrivals per hour during the period.",
        "required": "Required for aggregate uploads",
        "example": "30",
        "notes": "You do not need to calculate this if you upload customer events instead.",
    },
    {
        "field": "mu",
        "meaning": "Average customers one server can finish per hour.",
        "required": "Required for aggregate uploads",
        "example": "12",
        "notes": "You do not need to calculate this if you upload customer events instead.",
    },
    {
        "field": "c",
        "meaning": "How many servers served this line during the period. Always 1 for separate queues.",
        "required": "Required for aggregate uploads",
        "example": "3",
        "notes": "Must be a whole number of 1 or more.",
    },
    {
        "field": "variance",
        "meaning": "How much individual service times vary, in hours squared.",
        "required": "Required for separate queues; optional for a shared queue",
        "example": "0.006",
        "notes": "Best taken from measured service times. Must be zero or more.",
    },
    {
        "field": "arrival_time",
        "meaning": "The time the customer joined the waiting line.",
        "required": "Required for customer-event uploads",
        "example": "2026-01-05 09:14:32",
        "notes": "Any unambiguous date-time format works.",
    },
    {
        "field": "service_start",
        "meaning": "The time the cashier actually begins serving the customer.",
        "required": "Required for customer-event uploads",
        "example": "2026-01-05 09:16:05",
        "notes": "Cannot be before the arrival time.",
    },
    {
        "field": "service_end",
        "meaning": "The time the customer's service is completed.",
        "required": "Required for customer-event uploads",
        "example": "2026-01-05 09:21:40",
        "notes": "Cannot be before the service start. NovaQ calculates waiting time and service duration from these three timestamps.",
    },
    {
        "field": "K",
        "meaning": "Advanced: total customers allowed in the system including those in service.",
        "required": "Optional",
        "example": "15",
        "notes": "Only needed for finite-capacity studies configured in Setup.",
    },
    {
        "field": "theta",
        "meaning": "Advanced: how quickly waiting customers lose patience, per hour.",
        "required": "Optional",
        "example": "0.5",
        "notes": "Only needed when abandonment modeling is configured in Setup. Never used for separate queues.",
    },
    {
        "field": "server_cost",
        "meaning": "Advanced: cost of one server for one hour, used in optimization.",
        "required": "Optional",
        "example": "87",
        "notes": "Defaults apply when omitted.",
    },
]

# Clearly marked example rows. Every set validates cleanly on its own.
EXAMPLE_ROWS: dict[tuple[str, str], list[dict[str, object]]] = {
    ("shared_queue", "aggregate"): [
        {"time": "08:00-09:00", "lambda": 30.0, "mu": 12.0, "c": 3},
        {"time": "09:00-10:00", "lambda": 45.0, "mu": 12.0, "c": 4},
    ],
    ("separate_queues", "aggregate"): [
        {"time": "08:00-09:00", "queue_id": "queue_1", "lambda": 10.0, "mu": 12.0, "c": 1, "variance": 0.006},
        {"time": "08:00-09:00", "queue_id": "queue_2", "lambda": 10.0, "mu": 12.0, "c": 1, "variance": 0.006},
    ],
    ("shared_queue", "events"): [
        {"arrival_time": "2026-01-05 09:14:32", "service_start": "2026-01-05 09:16:05", "service_end": "2026-01-05 09:21:40"},
        {"arrival_time": "2026-01-05 09:15:10", "service_start": "2026-01-05 09:21:40", "service_end": "2026-01-05 09:26:02"},
    ],
    ("separate_queues", "events"): [
        {"arrival_time": "2026-01-05 09:14:32", "service_start": "2026-01-05 09:16:05", "service_end": "2026-01-05 09:21:40", "queue_id": "queue_1"},
        {"arrival_time": "2026-01-05 09:15:10", "service_start": "2026-01-05 09:21:40", "service_end": "2026-01-05 09:26:02", "queue_id": "queue_1"},
    ],
}

# Header aliases live in ingestion.py (single mapping engine) and are
# re-exported here so template/guide consumers share the exact contract.


def validate_template_key(structure: str, schema: str) -> tuple[str, str]:
    """Return the canonical (structure, schema) pair or raise ValueError."""
    if structure not in STRUCTURES:
        raise ValueError(f"Unknown queue structure: {structure}.")
    if schema not in SCHEMAS:
        raise ValueError(f"Unknown upload schema: {schema}.")
    return structure, schema


def template_columns(structure: str, schema: str) -> list[str]:
    """Canonical entry columns for one template, drawn from ingestion constants."""
    structure, schema = validate_template_key(structure, schema)
    columns = TEMPLATE_COLUMNS[(structure, schema)]
    if schema == "aggregate":
        assert set(REQUIRED_COLUMNS) <= set(columns), "aggregate template must cover required columns"
        if structure == "separate_queues":
            assert {"queue_id", "variance"} <= set(columns), "separate template must require queue identity and variance"
    else:
        assert EVENT_REQUIRED <= set(columns), "event template must cover event columns"
        if structure == "separate_queues":
            assert "queue_id" in columns, "separate event template must require queue identity"
    return list(columns)


def field_guide(structure: str, schema: str) -> list[dict[str, str]]:
    """Guide entries relevant to one template, entry columns first."""
    columns = template_columns(structure, schema)
    ordered = [e for c in columns for e in FIELD_GUIDE if e["field"] == c]
    rest = [e for e in FIELD_GUIDE if e["field"] not in columns]
    return ordered + rest


def example_rows(structure: str, schema: str) -> list[dict[str, object]]:
    """Marked example rows for one template."""
    validate_template_key(structure, schema)
    return [dict(row) for row in EXAMPLE_ROWS[(structure, schema)]]


def template_csv(structure: str, schema: str) -> bytes:
    """Clean CSV template: canonical headers, no data rows."""
    columns = template_columns(structure, schema)
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=columns)
    writer.writeheader()
    return buffer.getvalue().encode("utf-8")


def template_workbook(structure: str, schema: str) -> bytes:
    """XLSX template: Data Entry (headers only), Example, Field Guide."""
    columns = template_columns(structure, schema)
    guide = field_guide(structure, schema)
    examples = example_rows(structure, schema)
    workbook = Workbook()
    entry = workbook.active
    entry.title = "Data Entry"
    entry.append(columns)
    example = workbook.create_sheet("Example")
    example.append(columns)
    for row in examples:
        example.append([row.get(column) for column in columns])
    sheet = workbook.create_sheet("Field Guide")
    sheet.append(["Field", "What it means", "Required?", "Example", "Notes"])
    for item in guide:
        sheet.append([item["field"], item["meaning"], item["required"], item["example"], item["notes"]])
    for sheet in (entry, example):
        for index in range(1, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 20
    guide_sheet = workbook["Field Guide"]
    for index in range(1, 6):
        guide_sheet.column_dimensions[get_column_letter(index)].width = 32
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


__all__ = [
    "EXAMPLE_ROWS",
    "FIELD_GUIDE",
    "HEADER_ALIASES",
    "SCHEMAS",
    "STRUCTURES",
    "TEMPLATE_COLUMNS",
    "example_rows",
    "field_guide",
    "normalize_header_aliases",
    "template_columns",
    "template_csv",
    "template_workbook",
    "validate_template_key",
]
