"""Reports API coverage: PDF/Excel generation from datasets and scenarios."""

from __future__ import annotations

import io

import openpyxl

from tests.helpers import clear_cookies, create_user, csrf_header, login

CSV_GOOD = b"time,lambda,mu,c\n08:00-09:00,30,12,3\n09:00-10:00,45,12,4\n"


def make_dataset(client):
    response = client.post(
        "/datasets",
        headers=csrf_header(client),
        files={"file": ("segments.csv", CSV_GOOD, "application/octet-stream")},
    )
    return response.json()["dataset"]["id"]


def make_scenario(client):
    batch = client.post(
        "/optimize/batch",
        json={
            "segments": [
                {"time": "08:00-09:00", "lambda": 30, "mu": 12, "c": 3},
                {"time": "09:00-10:00", "lambda": 45, "mu": 12, "c": 4},
            ]
        },
    ).json()["results"]
    response = client.post(
        "/scenarios",
        headers=csrf_header(client),
        json={"name": "plan", "settings": {"target_utilization": 0.7}, "results": {"results": batch}},
    )
    return response.json()["scenario"]["id"]


def test_dataset_pdf_report(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    dataset_id = make_dataset(client)
    response = client.get(f"/reports/datasets/{dataset_id}/pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="novaq_datasets_{dataset_id}.pdf"'
    )
    assert response.content[:4] == b"%PDF"
    assert len(response.content) > 1000


def test_dataset_excel_report(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    dataset_id = make_dataset(client)
    response = client.get(f"/reports/datasets/{dataset_id}/excel")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.headers["content-disposition"] == (
        f'attachment; filename="novaq_datasets_{dataset_id}.xlsx"'
    )
    assert len(response.content) > 500

    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb["Segments"]
    assert [c.value for c in ws[1]] == [
        "time",
        "c_current",
        "rho_current",
        "rho_current_status",
        "Wq_current",
        "metric_provenance",
        "selected_model",
    ]
    assert ws["A2"].value == "08:00-09:00"
    assert ws["F2"].value == "analytical"
    assert ws["B2"].value == 3
    assert 0 < ws["C2"].value < 1
    assert ws["D2"].value in {"Lean", "Normal", "Peak", "Critical", "Unstable"}
    assert 0 < ws["E2"].value < 60
    summary_labels = [wb["Summary"].cell(row=r, column=1).value for r in range(1, 8)]
    assert "Avg Wait Current, modeled (min)" in summary_labels
    assert "Avg Utilization Current" in summary_labels


def test_scenario_pdf_report(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    scenario_id = make_scenario(client)
    response = client.get(f"/reports/scenarios/{scenario_id}/pdf")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="novaq_scenarios_{scenario_id}.pdf"'
    )
    assert response.content[:4] == b"%PDF"


def test_scenario_excel_report(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    scenario_id = make_scenario(client)
    response = client.get(f"/reports/scenarios/{scenario_id}/excel")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert response.headers["content-disposition"] == (
        f'attachment; filename="novaq_scenarios_{scenario_id}.xlsx"'
    )


def test_reports_require_auth(client):
    assert client.get("/reports/datasets/1/pdf").status_code == 401
    assert client.get("/reports/scenarios/1/pdf").status_code == 401


def test_dataset_report_missing_404(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    assert client.get("/reports/datasets/99999/pdf").status_code == 404


def test_scenario_report_missing_404(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    assert client.get("/reports/scenarios/99999/pdf").status_code == 404


def test_cross_user_dataset_report_404(db_engine, client):
    create_user(db_engine, "alice@example.com", "pw")
    login(client, "alice@example.com", "pw")
    dataset_id = make_dataset(client)
    clear_cookies(client)
    create_user(db_engine, "bob@example.com", "pw")
    login(client, "bob@example.com", "pw")
    assert client.get(f"/reports/datasets/{dataset_id}/pdf").status_code == 404


def test_cross_user_scenario_report_404(db_engine, client):
    create_user(db_engine, "alice@example.com", "pw")
    login(client, "alice@example.com", "pw")
    scenario_id = make_scenario(client)
    clear_cookies(client)
    create_user(db_engine, "bob@example.com", "pw")
    login(client, "bob@example.com", "pw")
    assert client.get(f"/reports/scenarios/{scenario_id}/excel").status_code == 404


def test_scenario_without_comparison_results_422(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    response = client.post(
        "/scenarios",
        headers=csrf_header(client),
        json={"name": "bare", "settings": {}, "results": {"note": "nothing here"}},
    )
    scenario_id = response.json()["scenario"]["id"]
    assert client.get(f"/reports/scenarios/{scenario_id}/pdf").status_code == 422


def test_blocked_scenario_reports_stay_unavailable(db_engine, client):
    """Honest INVALID_INPUT rows export as 200s whose evidence stays N/A/incomplete."""
    create_user(db_engine, "blocked@example.com", "pw")
    login(client, "blocked@example.com", "pw")
    batch = client.post(
        "/optimize/batch",
        json={"segments": [{
            "time": "08:00", "lambda": 5.0, "mu": 4.0, "c": 1, "variance": 0.0025,
            "queue_structure": "separate_queues", "model_id": "parallel_mg1",
        }]},
    ).json()["results"]
    assert batch[0]["c_optimal"] is None
    scenario_id = client.post(
        "/scenarios",
        headers=csrf_header(client),
        json={"name": "blocked", "settings": {"target_utilization": 0.7}, "results": {"results": batch}},
    ).json()["scenario"]["id"]

    pdf = client.get(f"/reports/scenarios/{scenario_id}/pdf")
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"

    excel = client.get(f"/reports/scenarios/{scenario_id}/excel")
    assert excel.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(excel.content))
    cells = {wb["Summary"].cell(r, 1).value: wb["Summary"].cell(r, 2).value for r in range(1, 25)}
    assert cells["Total Savings"] == "N/A"
    assert cells["Total Server Change"] == "N/A"
    recs = [wb["Recommendations"].cell(r, 1).value for r in range(2, 6)]
    assert any(isinstance(rec, str) and "Comparison incomplete" in rec for rec in recs)


def test_unknown_format_404(db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    dataset_id = make_dataset(client)
    assert client.get(f"/reports/datasets/{dataset_id}/csv").status_code == 404


def test_analysis_scoped_report_uses_decision_recommendation(monkeypatch, db_engine, client):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    scenario_id = make_scenario(client)
    scenario = client.get(f"/scenarios/{scenario_id}").json()["scenario"]
    analysis_id = scenario["analysis_id"]

    captured: dict[str, list[str]] = {}

    def fake_pdf(current, recommended, frame, recommendations):
        captured["recommendations"] = recommendations
        return io.BytesIO(b"%PDF decision")

    monkeypatch.setattr("backend.api.reports.generate_pdf_report", fake_pdf)
    response = client.get(
        f"/reports/scenarios/{scenario_id}/pdf?analysis_id={analysis_id}"
    )
    assert response.status_code == 200
    assert captured["recommendations"] == [
        "Management recommendation unavailable: complete Decision for the "
        "currently selected and simulated Scenario."
    ]


def test_analysis_scoped_reports_consume_the_persisted_decision(
    monkeypatch, db_engine, client
):
    create_user(db_engine, "u@example.com", "pw")
    login(client, "u@example.com", "pw")
    scenario_id = make_scenario(client)
    scenario = client.get(f"/scenarios/{scenario_id}").json()["scenario"]
    analysis_id = scenario["analysis_id"]
    decision = {
        "headline": "Adopt Scenario A.",
        "recommendation": "All validation intervals passed.",
        "rationale": ["Modeled cost is lower."],
        "provenance_warning": "Simulation is decision support.",
    }
    monkeypatch.setattr(
        "backend.api.reports.current_decision_for_report",
        lambda *args: decision,
    )

    captured: dict[str, list[str]] = {}

    def fake_pdf(current, recommended, frame, recommendations):
        captured["recommendations"] = recommendations
        return io.BytesIO(b"%PDF decision")

    monkeypatch.setattr("backend.api.reports.generate_pdf_report", fake_pdf)
    pdf = client.get(
        f"/reports/scenarios/{scenario_id}/pdf?analysis_id={analysis_id}"
    )
    assert pdf.status_code == 200
    assert captured["recommendations"] == [
        "Adopt Scenario A.",
        "All validation intervals passed.",
        "Modeled cost is lower.",
        "Simulation is decision support.",
    ]

    excel = client.get(
        f"/reports/scenarios/{scenario_id}/excel?analysis_id={analysis_id}"
    )
    assert excel.status_code == 200
    workbook = openpyxl.load_workbook(io.BytesIO(excel.content))
    assert [cell.value for cell in workbook["Recommendations"]["A"]] == [
        "Recommendation Evidence",
        "Adopt Scenario A.",
        "All validation intervals passed.",
        "Modeled cost is lower.",
        "Simulation is decision support.",
    ]
