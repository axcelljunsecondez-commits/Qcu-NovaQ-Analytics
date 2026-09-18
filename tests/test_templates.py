"""Template/example/guide parity with the ingestion contract, plus header aliases."""

from __future__ import annotations

import io

import pandas as pd
import pytest

from backend.api.analysis_schemas import QueueSetup, QueueStructure
from backend.data.analysis_ingestion import EVENT_REQUIRED, AnalysisIngestionError, normalize_analysis_input
from backend.data.ingestion import REQUIRED_COLUMNS, validate_and_normalize
from backend.data.templates import (
    HEADER_ALIASES,
    example_rows,
    field_guide,
    normalize_header_aliases,
    template_columns,
    template_csv,
    template_workbook,
)


def _shared_setup(**overrides):
    values = {
        "queue_structure": QueueStructure.shared_queue,
        "fixed_server_count": None,
        "staffing_varies_by_period": False,
        "capacity_mode": "unlimited",
        "total_system_capacity": None,
        "abandonment_mode": "not_modeled",
        "patience_rate_per_hour": None,
        "segments": [],
        "separate_queue_closure_policy": "drain_existing",
        "queue_ids": [],
    }
    values.update(overrides)
    return QueueSetup(**values)


def _separate_setup(**overrides):
    values = {
        "queue_structure": QueueStructure.separate_queues,
        "fixed_server_count": 1,
        "staffing_varies_by_period": False,
        "capacity_mode": "unlimited",
        "total_system_capacity": None,
        "abandonment_mode": "not_modeled",
        "patience_rate_per_hour": None,
        "segments": [],
        "separate_queue_closure_policy": "drain_existing",
        "queue_ids": ["queue_1", "queue_2"],
    }
    values.update(overrides)
    return QueueSetup(**values)


@pytest.mark.parametrize(
    "structure,schema",
    [
        ("shared_queue", "aggregate"),
        ("separate_queues", "aggregate"),
        ("shared_queue", "events"),
        ("separate_queues", "events"),
    ],
)
def test_template_columns_cover_ingestion_contract(structure, schema):
    columns = template_columns(structure, schema)
    if schema == "aggregate":
        assert set(REQUIRED_COLUMNS) <= set(columns)
        if structure == "separate_queues":
            assert {"queue_id", "variance"} <= set(columns)
    else:
        assert EVENT_REQUIRED <= set(columns)
        if structure == "separate_queues":
            assert "queue_id" in columns


@pytest.mark.parametrize(
    "structure,schema",
    [
        ("shared_queue", "aggregate"),
        ("separate_queues", "aggregate"),
        ("shared_queue", "events"),
        ("separate_queues", "events"),
    ],
)
def test_example_rows_validate_cleanly(structure, schema):
    columns = template_columns(structure, schema)
    frame = pd.DataFrame(example_rows(structure, schema), columns=columns)
    if schema == "aggregate" and structure == "shared_queue":
        setup: QueueSetup = _shared_setup()
    elif schema == "aggregate":
        setup = _separate_setup()
    else:
        setup = (
            _separate_setup(fixed_server_count=1)
            if structure == "separate_queues"
            else _shared_setup(fixed_server_count=2)
        )
    records, provenance = normalize_analysis_input(frame, setup)
    if schema == "aggregate":
        assert len(records) == len(example_rows(structure, schema))
    else:
        # Same-hour events aggregate into one hourly record by design.
        assert len(records) == 1
        assert records[0]["lambda"] == 2.0
    assert provenance["schema"] == ("customer_events" if schema == "events" else "aggregate")
    assert provenance["column_mapping"] == {}


def test_csv_template_round_trip():
    content = template_csv("separate_queues", "aggregate").decode("utf-8")
    header = content.splitlines()[0]
    assert header == "time,queue_id,lambda,mu,c,variance"
    frame = pd.read_csv(io.StringIO(content + "08:00-09:00,queue_1,10,12,1,0.006\n"))
    ok, message, _ = validate_and_normalize(frame)
    assert ok, message


def test_xlsx_template_first_sheet_round_trip():
    data = template_workbook("separate_queues", "aggregate")
    frame = pd.read_excel(io.BytesIO(data), engine="openpyxl")
    assert list(frame.columns) == ["time", "queue_id", "lambda", "mu", "c", "variance"]
    assert frame.empty


def test_xlsx_template_has_example_and_guide_sheets():
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(template_workbook("shared_queue", "events")))
    assert workbook.sheetnames == ["Data Entry", "Example", "Field Guide"]
    guide_fields = [row[0].value for row in workbook["Field Guide"].iter_rows(min_row=2)]
    for required in ("arrival_time", "service_start", "service_end"):
        assert required in guide_fields


def test_field_guide_covers_every_template_column():
    for structure in ("shared_queue", "separate_queues"):
        for schema in ("aggregate", "events"):
            guided = {entry["field"] for entry in field_guide(structure, schema)}
            assert set(template_columns(structure, schema)) <= guided


def test_business_headers_map_to_canonical_fields():
    frame = pd.DataFrame(
        [
            {"Period": "08:00-09:00", "Lane": "queue_1", "Arrival rate": 10.0,
             "Service rate": 12.0, "Servers": 1, "Service variance": 0.006},
        ]
    )
    records, provenance = normalize_analysis_input(frame, _separate_setup())
    assert len(records) == 1
    assert records[0]["queue_id"] == "queue_1"
    assert records[0]["lambda"] == 10.0
    assert provenance["column_mapping"]


def test_canonical_headers_win_over_aliases():
    renamed, mapping = normalize_header_aliases(["time", "Period"])
    assert renamed == ["time", "Period"]
    assert mapping == {}


def test_missing_event_field_error_stays_explanatory():
    frame = pd.DataFrame(
        [{"arrival_time": "2026-01-05 09:14:32", "service_end": "2026-01-05 09:21:40"}]
    )
    with pytest.raises(AnalysisIngestionError, match="service_start.*begins serving"):
        normalize_analysis_input(frame, _shared_setup(fixed_server_count=2))


def test_missing_values_are_not_zeroed():
    frame = pd.DataFrame([{"time": "08:00-09:00", "lambda": 30.0, "mu": 12.0, "c": 3}])
    ok, _, normalized = validate_and_normalize(frame)
    assert ok
    assert "variance" not in normalized.columns


def test_template_endpoints(db_engine, client):
    from tests.helpers import create_user, csrf_header, login

    create_user(db_engine, "template-owner@example.com", "pw")
    login(client, "template-owner@example.com", "pw")
    guide = client.get("/templates/guide?structure=separate_queues&schema=aggregate")
    assert guide.status_code == 200, guide.text
    assert guide.json()["columns"] == ["time", "queue_id", "lambda", "mu", "c", "variance"]
    download = client.get("/templates/download?structure=separate_queues&schema=aggregate&format=csv")
    assert download.status_code == 200
    assert download.headers["content-type"].startswith("text/csv")
    xlsx = client.get("/templates/download?structure=shared_queue&schema=events&format=xlsx")
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )
    bad = client.get("/templates/guide?structure=pooled&schema=aggregate")
    assert bad.status_code == 422
    assert HEADER_ALIASES["queue_id"]


def test_separate_events_workbook_has_events_staff_and_breaks_sheets():
    import openpyxl

    from backend.data import uploads

    data = template_workbook("separate_queues", "events")
    workbook = openpyxl.load_workbook(io.BytesIO(data))
    assert workbook.sheetnames == ["events", "staff", "breaks", "Example", "Field Guide"]
    headers = {name: [cell.value for cell in workbook[name][1]] for name in ("events", "staff", "breaks")}
    assert headers == {
        "events": ["arrival_time", "service_start", "service_end", "queue_id"],
        "staff": ["queue_id", "shift_start", "shift_end"],
        "breaks": ["queue_id", "break_name", "start", "minutes"],
    }
    guide_fields = [row[0].value for row in workbook["Field Guide"].iter_rows(min_row=2)]
    for field in ("shift_start", "shift_end", "break_name", "start", "minutes"):
        assert field in guide_fields
    sheets = uploads.parse_upload_workbook("template.xlsx", data)
    assert sheets is not None and sheets["staff"] is None and sheets["breaks"] is None
    # The shared template and its guide are unchanged.
    shared = openpyxl.load_workbook(io.BytesIO(template_workbook("shared_queue", "events")))
    assert shared.sheetnames == ["Data Entry", "Example", "Field Guide"]
    assert "shift_start" not in [row[0].value for row in shared["Field Guide"].iter_rows(min_row=2)]
