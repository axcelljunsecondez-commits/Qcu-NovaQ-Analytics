"""Full Separate-Queue report: normalized model from persisted evidence only.

The builder normalizes the exact selected chain (analysis, dataset, scenario,
schedule, DES, MC, validation, decision, current) into one model shared by
preview, PDF, and Excel. It never reruns engines, never recalculates an
optimum, and never fabricates savings, ROI, or ADOPT.
"""
from __future__ import annotations

import pytest

from backend.api.scenarios import setup_fingerprint
from backend.reports.separate_report import build_separate_report_model


def _chain(**over):
    base = {
        "analysis": {"id": 7, "name": "East lanes",
                     "queue_structure": "separate_queues",
                     "queue_ids": ["east-07", "lane-A"],
                     "closure_policy": "drain_existing"},
        "dataset": {"id": 2, "name": "lanes.csv", "row_count": 2,
                    "periods": ["08:00"], "source_format": "csv",
                    "validation_ok": True},
        "scenario": {"id": 9, "name": "Optimal @ 70%", "target": 0.70,
                     "engine_version": "novaq-2026-09-separate-des-v1",
                     "calculated_at": "2026-09-18T00:00:00+00:00",
                     "dataset_id": 2},
        "schedule": {
            "target_utilization": 0.70,
            "evaluation_method": "DES_REPLICATIONS",
            "des": {"replications": 5, "base_seed": 42,
                    "duration_hours": 8.0, "max_events": 5000},
            "periods": [{
                "time": "08:00", "current_active_lanes": ["east-07", "lane-A"],
                "optimal_active_lanes": 1, "adjustment": -1,
                "optimum": {"active_queue_ids": ["east-07"],
                            "inactive_queue_ids": ["lane-A"],
                            "candidate_utilization": 0.55,
                            "total_cost": 180.0},
            }],
        },
        "current": {
            "periods": [{
                "time": "08:00",
                "queues": [
                    {"queue_id": "east-07", "lambda": 4.0, "mu": 11.0,
                     "rho": 0.36, "Wq": 0.05, "model": "M/G/1",
                     "stable": True},
                    {"queue_id": "lane-A", "lambda": 2.0, "mu": 11.0,
                     "rho": 0.18, "Wq": 0.02, "model": "M/G/1",
                     "stable": True},
                ],
            }],
            "kpis": {"avg_waiting_time": 0.04, "max_utilization": 0.36,
                     "total_waiting_cost": 25.0},
            "models": ["M/G/1"],
        },
        "des": {
            "job_id": 10, "seed": 7, "duration_hours": 8.0,
            "overall_conservation": True, "overall_status": "COMPLETED",
            "arrival_method": "single conserved Poisson stream at total lambda",
            "routing_policy": "shortest system size with seeded fair ties",
            "service_sampling_method": "empirical per-lane resampling",
            "execution": "selected-plan routing DES (period-independent, no carryover)",
            "periods": [{
                "time": "08:00", "active_queue_ids": ["east-07"],
                "conservation": True,
                "lanes": [
                    {"queue_id": "east-07", "arrivals": 46, "served": 45,
                     "waiting": 1, "Wq": 0.09, "rho": 0.55, "max_queue": 4,
                     "active": True},
                    {"queue_id": "lane-A", "arrivals": 0, "served": 0,
                     "waiting": 0, "Wq": None, "rho": 0.0, "max_queue": 0,
                     "active": False},
                ],
                "trace": {"event_count": 120, "truncated": False,
                          "queue_ids": ["east-07"]},
            }],
        },
        "mc": {
            "job_id": 11, "num_trials": 2000, "failure_threshold": 0.75,
            "failure_rate_cap": 0.05,
            "method": "analytical parameter perturbation at DES-measured lane loads",
            "lanes": [{
                "time": "08:00", "queue_id": "east-07", "lambda": 5.75,
                "failure_rate": 0.0, "failure_rate_ci": [0.0, 0.002],
                "adequate": True, "status": "PASS",
            }],
        },
        "validation": {
            "job_id": 12,
            "verdict": "pass",
            "periods": [{
                "time": "08:00", "active_queue_ids": ["east-07"],
                "status": "pass",
                "queues": [{
                    "queue_id": "east-07", "rho_sim": 0.55, "Wq_sim": 0.09,
                    "served": 45, "mc_failure_rate": 0.0,
                    "validation_verdict": "pass",
                }],
            }],
        },
        "decision": {
            "job_id": 13, "status": "conditional",
            "headline": 'Consider Scenario "Optimal @ 70%" conditionally.',
            "recommendation": "All 1 validation checks passed, but modeled savings are unavailable.",
            "rationale": ["Selected Scenario: Optimal @ 70% (ID 9)."],
            "facts": {"lane_delta": -1, "validation_checks": 1,
                      "failed_checks": 0},
            "failed_periods": [],
        },
        "comparison_plans": [
            {"scenario_id": 9, "name": "Optimal @ 70%", "target": 0.70,
             "overall": "COMPLETE"},
            {"scenario_id": 10, "name": "Optimal @ 60%", "target": 0.60,
             "overall": "COMPLETE"},
        ],
    }
    base.update(over)
    return base


def test_model_preserves_evidence_without_fabrication():
    model = build_separate_report_model(_chain())
    assert model["decision"]["status"] == "conditional"
    assert model["decision"]["headline"].startswith("Consider")
    assert model["schedule"]["periods"][0]["optimal_active_lanes"] == 1
    assert model["schedule"]["periods"][0]["optimum"]["active_queue_ids"] == ["east-07"]
    assert model["current"]["periods"][0]["queues"][0]["queue_id"] == "east-07"
    assert model["des"]["periods"][0]["lanes"][1]["Wq"] is None
    assert model["cost"]["savings"] is None
    assert model["cost"]["roi"] is None
    assert model["cost"]["current_total"] is None
    assert model["cost"]["current_total_reason"] is not None
    assert model["cost"]["selected_total"] == 180.0
    assert model["mc"]["lanes"][0]["failure_rate"] == 0.0
    assert model["validation"]["verdict"] == "pass"
    assert model["provenance"]["scenario_id"] == 9
    assert model["provenance"]["des_job_id"] == 10
    assert model["provenance"]["mc_job_id"] == 11
    assert model["provenance"]["validation_job_id"] == 12
    assert model["provenance"]["decision_job_id"] == 13
    assert "period-independent" in " ".join(model["limitations"])
    assert "ADOPT" in " ".join(model["limitations"])
    assert model["staffing_title"] == "Selected Staffing Schedule"


def test_revise_model_titles_schedule_without_approval_language():
    chain = _chain()
    chain["decision"] = {**chain["decision"], "status": "revise",
                         "headline": 'Do not adopt Scenario "Optimal @ 70%" yet.'}
    model = build_separate_report_model(chain)
    assert model["decision"]["status"] == "revise"
    assert model["staffing_title"] == "Selected Plan Requiring Revision"
    assert "Approved" not in model["staffing_title"]


def test_pdf_sections_match_preview_model_values():
    from backend.reports.report_export import (
        generate_separate_pdf_report,
        separate_pdf_section_headings,
    )

    model = build_separate_report_model(_chain())
    headings = separate_pdf_section_headings(model)
    assert "Decision: CONDITIONAL" in headings
    assert "Limitations" in headings
    assert "Selected Staffing Schedule" in headings
    buffer = generate_separate_pdf_report(model)
    assert buffer.getvalue()[:5] == b"%PDF-"
    assert len(buffer.getvalue()) > 2000


def test_excel_roundtrip_preserves_types_and_nas():
    import openpyxl

    from backend.reports.report_export import generate_separate_excel_report

    model = build_separate_report_model(_chain())
    buffer = generate_separate_excel_report(model)
    workbook = openpyxl.load_workbook(buffer)
    assert "Overview" in workbook.sheetnames
    assert "DES" in workbook.sheetnames
    assert "Provenance" in workbook.sheetnames
    overview = {row[0].value: row[1].value
                for row in workbook["Overview"].iter_rows(min_row=2)}
    assert overview["Decision"] == "conditional"
    assert overview["Scenario"] == "Optimal @ 70%"
    des = workbook["DES"]
    header = [cell.value for cell in des[1]]
    rho_col = header.index("Utilization") + 1
    wq_col = header.index("Wq (min)") + 1
    peak_col = header.index("Peak Mean Utilization") + 1
    assert isinstance(des.cell(row=2, column=rho_col).value, float)
    assert des.cell(row=2, column=peak_col).value == pytest.approx(0.55)
    assert des.cell(row=2, column=rho_col).number_format == "0.0%"
    assert des.cell(row=2, column=wq_col).value == pytest.approx(5.4)
    assert workbook["Cost"]["B3"].value == "N/A"
    sched = workbook["Selected Staffing"]
    staffing_col = [cell.value for cell in sched[1]].index("Selected Active") + 1
    assert sched.cell(row=2, column=staffing_col).value == 1
    assert isinstance(sched.cell(row=2, column=staffing_col).value, int)
    assert workbook["Provenance"]["B2"].value == 9
    limitations = [row[0].value for row in workbook["Limitations"].iter_rows(min_row=2)]
    assert any("period-independent" in str(line) for line in limitations)
    assert not any("continuous" in str(line).lower() for line in limitations)


# --- API layer -----------------------------------------------------------------------

def _api_workspace(db_engine, email):
    from backend.db.models import AnalysisProject, Dataset, Scenario, User
    from tests.helpers import create_user, make_sessionmaker

    samples = [0.05, 0.08, 0.10, 0.12]
    rows = [
        {"time": "08:00", "segment_id": "s1", "queue_id": "east-07",
         "lambda": 2.0, "mu": 11.0, "c": 1, "variance": 0.0009,
         "service_samples_hours": list(samples)},
        {"time": "08:00", "segment_id": "s1", "queue_id": "lane-A",
         "lambda": 1.0, "mu": 11.0, "c": 1, "variance": 0.0009,
         "service_samples_hours": list(samples)},
    ]
    setup = {
        "queue_structure": "separate_queues", "fixed_server_count": 1,
        "staffing_varies_by_period": False, "capacity_mode": "unlimited",
        "total_system_capacity": None, "abandonment_mode": "not_modeled",
        "patience_rate_per_hour": None,
        "segments": [{"id": "s1", "start_time": "08:00:00",
                      "end_time": "09:00:00", "active_queue_ids": None}],
        "separate_queue_closure_policy": "drain_existing",
        "queue_ids": ["east-07", "lane-A"],
    }
    options = {
        "target_utilization": 0.70, "server_cost_per_hr": 87.0,
        "customer_waiting_cost": 100.0, "min_active_lanes": None,
        "max_active_lanes": None, "lambda_multiplier": 1.0,
        "des": {"replications": 2, "base_seed": 101,
                "duration_hours": 4.0, "max_events": 500},
    }
    period = {
        "time": "08:00", "overall": "OPTIMAL", "reason": None,
        "current_active_lanes": ["east-07", "lane-A"],
        "optimal_active_lanes": 1, "adjustment": -1,
        "optimum": {
            "active_lane_count": 1, "recommendation": "REDUCE",
            "total_cost": 120.0, "server_cost": 87.0, "waiting_cost": 33.0,
            "candidate_utilization": 0.45, "estimated_optimal": True,
            "active_queue_ids": ["east-07"], "inactive_queue_ids": ["lane-A"],
        },
        "candidates": [], "evaluation_method": "DES_REPLICATIONS",
        "replication_seeds": [101, 102], "target_utilization": 0.70,
    }
    schedule = {"overall": "COMPLETE", "reason": None,
                "target_utilization": 0.70,
                "evaluation_method": "DES_REPLICATIONS", "periods": [period],
                "des": dict(options["des"])}
    user = create_user(db_engine, email, "pw")
    with make_sessionmaker(db_engine)() as db:
        analysis = AnalysisProject(
            user_id=user.id, name="East lanes",
            queue_setup_json=setup, setup_status="ready")
        db.add(analysis)
        db.flush()
        dataset = Dataset(
            user_id=user.id, analysis_id=analysis.id, name="Lanes",
            source_filename="lanes.csv", source_format="csv", row_count=2,
            normalized_json=rows,
            validation_report_json={"ok": True, "message": "ok"})
        db.add(dataset)
        db.flush()
        scenario = Scenario(
            user_id=user.id, analysis_id=analysis.id, dataset_id=dataset.id,
            name="Optimal @ 70%",
            settings_json={**options, "calculation": {
                "schema_version": 2,
                "engine_version": "novaq-2026-09-separate-des-v1",
                "dataset_id": dataset.id, "dataset_row_count": 2,
                "options": options, "calculated_at": "2026-09-18T00:00:00+00:00",
                "setup_hash": setup_fingerprint(setup)}},
            results_json={"schedule": schedule})
        db.add(scenario)
        db.commit()
        return analysis.id, dataset.id, scenario.id


def _full_chain(client, analysis_id, headers, scenario_id):
    assert client.post(
        f"/analyses/{analysis_id}/workflow/selection", headers=headers,
        json={"scenario_id": scenario_id}).status_code == 200
    assert client.post(
        f"/analyses/{analysis_id}/workflow/simulation/des/selected",
        headers=headers, json={"seed": 7}).status_code == 200
    assert client.post(
        f"/analyses/{analysis_id}/workflow/simulation/mc/selected",
        headers=headers, json={}).status_code == 200
    assert client.post(
        f"/analyses/{analysis_id}/workflow/simulation/validation/selected",
        headers=headers, json={}).status_code == 200
    decided = client.post(
        f"/analyses/{analysis_id}/workflow/decision/selected",
        headers=headers, json={})
    assert decided.status_code == 200, decided.text


def test_report_gate_blocks_without_decision_chain(db_engine, client):
    from tests.helpers import csrf_header, login

    analysis_id, _, _ = _api_workspace(db_engine, "u@example.com")
    login(client, "u@example.com", "pw")
    headers = csrf_header(client)
    assert client.get(
        f"/reports/analyses/{analysis_id}/selected/preview",
        headers=headers).status_code == 409


def test_full_report_preview_pdf_excel_share_one_model(db_engine, client):
    from tests.helpers import csrf_header, login

    analysis_id, _, scenario_id = _api_workspace(db_engine, "v@example.com")
    login(client, "v@example.com", "pw")
    headers = csrf_header(client)
    _full_chain(client, analysis_id, headers, scenario_id)
    from backend.db.models import Job, Scenario
    from tests.helpers import make_sessionmaker

    with make_sessionmaker(db_engine)() as db:
        before = {row.id: dict(row.result_json) for row in db.query(Job).all()}
        scenario_before = dict(db.get(Scenario, scenario_id).results_json)
    preview = client.get(
        f"/reports/analyses/{analysis_id}/selected/preview", headers=headers)
    assert preview.status_code == 200, preview.text
    model = preview.json()["model"]
    assert model["provenance"]["scenario_id"] == scenario_id
    assert model["decision"]["status"] == "conditional"
    assert model["cost"]["savings"] is None
    assert model["cost"]["roi"] is None
    assert model["cost"]["current_total"] is None
    assert model["schedule"]["periods"][0]["optimal_active_lanes"] == 1
    pdf = client.get(
        f"/reports/analyses/{analysis_id}/selected/pdf", headers=headers)
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content[:5] == b"%PDF-"
    excel = client.get(
        f"/reports/analyses/{analysis_id}/selected/excel", headers=headers)
    assert excel.status_code == 200
    assert "spreadsheetml" in excel.headers["content-type"]
    import openpyxl

    workbook = openpyxl.load_workbook(filename=__import__("io").BytesIO(excel.content))
    overview = {row[0].value: row[1].value
                for row in workbook["Overview"].iter_rows(min_row=2)}
    assert overview["Decision"] == model["decision"]["status"]
    assert overview["Scenario"] == model["overview"]["scenario_name"]
    assert workbook["Cost"]["B3"].value == "N/A"
    with make_sessionmaker(db_engine)() as db:
        for row in db.query(Job).all():
            assert dict(row.result_json) == before[row.id]
        assert db.get(Scenario, scenario_id).results_json == scenario_before


def test_stale_chain_blocks_report_and_scenario_change_rescopes(db_engine, client):
    from backend.db.models import Dataset, User
    from tests.helpers import csrf_header, login, make_sessionmaker

    analysis_id, _, scenario_id = _api_workspace(db_engine, "u@example.com")
    login(client, "u@example.com", "pw")
    headers = csrf_header(client)
    _full_chain(client, analysis_id, headers, scenario_id)
    with make_sessionmaker(db_engine)() as db:
        user = db.query(User).filter_by(email="u@example.com").one()
        newer = Dataset(
            user_id=user.id, analysis_id=analysis_id, name="Recut",
            source_filename="r.csv", source_format="csv", row_count=2,
            normalized_json=[
                {"time": "08:00", "segment_id": "s1", "queue_id": "east-07",
                 "lambda": 2.0, "mu": 11.0, "c": 1, "variance": 0.0009,
                 "service_samples_hours": [0.05, 0.08, 0.10, 0.12]},
                {"time": "08:00", "segment_id": "s1", "queue_id": "lane-A",
                 "lambda": 1.0, "mu": 11.0, "c": 1, "variance": 0.0009,
                 "service_samples_hours": [0.05, 0.08, 0.10, 0.12]},
            ],
            validation_report_json={"ok": True, "message": "ok"})
        db.add(newer)
        db.commit()
    blocked = client.get(
        f"/reports/analyses/{analysis_id}/selected/preview", headers=headers)
    assert blocked.status_code in (404, 409, 422)


def test_report_generation_makes_no_engine_calls_and_leaves_evidence_untouched():
    import inspect

    import backend.api.reports as reports_api

    source = inspect.getsource(reports_api)
    for forbidden in ("optimize_separate", "_optimize_separate_des",
                      "rank_des_candidates", "evaluate_candidate_with_des",
                      "mc_simulate_segments", "_run_routing_des",
                      "validate_selected_plan", "_derive_selected_decision",
                      "run_selected_validation", "create_selected_decision"):
        assert forbidden not in source


# --- hardening: staleness / identity / nulls / no-engine -----------------------------

from backend.db.models import AnalysisProject, Job, Scenario, User
from tests.helpers import csrf_header, login, make_sessionmaker


def _add_scenario(db_engine, email, analysis_id, dataset_id, name, target):
    with make_sessionmaker(db_engine)() as db:
        user = db.query(User).filter_by(email=email).one()
        options = {
            "target_utilization": target, "server_cost_per_hr": 87.0,
            "customer_waiting_cost": 100.0, "min_active_lanes": None,
            "max_active_lanes": None, "lambda_multiplier": 1.0,
            "des": {"replications": 2, "base_seed": 101,
                    "duration_hours": 4.0, "max_events": 500},
        }
        period = {
            "time": "08:00", "overall": "OPTIMAL", "reason": None,
            "current_active_lanes": ["east-07", "lane-A"],
            "optimal_active_lanes": 1, "adjustment": -1,
            "optimum": {
                "active_lane_count": 1, "recommendation": "REDUCE",
                "total_cost": 120.0, "server_cost": 87.0, "waiting_cost": 33.0,
                "candidate_utilization": 0.45, "estimated_optimal": True,
                "active_queue_ids": ["east-07"], "inactive_queue_ids": ["lane-A"],
            },
            "candidates": [], "evaluation_method": "DES_REPLICATIONS",
            "replication_seeds": [101, 102], "target_utilization": target,
        }
        schedule = {"overall": "COMPLETE", "reason": None,
                    "target_utilization": target,
                    "evaluation_method": "DES_REPLICATIONS", "periods": [period],
                    "des": dict(options["des"])}
        scenario = Scenario(
            user_id=user.id, analysis_id=analysis_id, dataset_id=dataset_id,
            name=name,
            settings_json={**options, "calculation": {
                "schema_version": 2,
                "engine_version": "novaq-2026-09-separate-des-v1",
                "dataset_id": dataset_id, "dataset_row_count": 2,
                "options": options, "calculated_at": "2026-09-18T00:00:00+00:00",
                "setup_hash": setup_fingerprint(
                    db.get(AnalysisProject, analysis_id).queue_setup_json)}},
            results_json={"schedule": schedule})
        db.add(scenario)
        db.commit()
        return scenario.id


def _insert_job(db_engine, email, kind, analysis_id, scenario_id, params, result):
    from datetime import datetime, timezone

    with make_sessionmaker(db_engine)() as db:
        user = db.query(User).filter_by(email=email).one()
        job = Job(
            user_id=user.id, kind=kind, status="completed",
            params_json={"analysis_id": analysis_id, "scenario_id": scenario_id,
                         "setup_hash": setup_fingerprint(
                             db.get(AnalysisProject, analysis_id).queue_setup_json),
                         **params},
            result_json=result, tenant_id=user.tenant_id,
            finished_at=datetime.now(timezone.utc))
        db.add(job)
        db.commit()
        return job.id


def _preview(client, analysis_id, headers):
    return client.get(
        f"/reports/analyses/{analysis_id}/selected/preview", headers=headers)


def _chain_ids(client, analysis_id, headers, scenario_id):
    _full_chain(client, analysis_id, headers, scenario_id)
    workflow = client.get(f"/analyses/{analysis_id}/workflow",
                          headers=headers).json()
    return {
        "des": workflow["des"]["id"],
        "mc": workflow["mc"]["id"],
        "validation": workflow["validation"]["id"],
        "decision": workflow["decision"]["id"],
    }


def test_two_targets_stay_pinned_to_selected_chain(db_engine, client):
    analysis_id, dataset_id, id_a = _api_workspace(db_engine, "u@example.com")
    login(client, "u@example.com", "pw")
    headers = csrf_header(client)
    ids_a = _chain_ids(client, analysis_id, headers, id_a)
    id_b = _add_scenario(db_engine, "u@example.com", analysis_id, dataset_id,
                         "Optimal @ 80%", 0.80)
    ids_b = _chain_ids(client, analysis_id, headers, id_b)
    assert ids_a != ids_b
    assert client.post(
        f"/analyses/{analysis_id}/workflow/selection", headers=headers,
        json={"scenario_id": id_a}).status_code == 200
    model = _preview(client, analysis_id, headers).json()["model"]
    provenance = model["provenance"]
    assert model["overview"]["target"] == 0.70
    assert provenance["scenario_id"] == id_a
    assert provenance["des_job_id"] == ids_a["des"]
    assert provenance["mc_job_id"] == ids_a["mc"]
    assert provenance["validation_job_id"] == ids_a["validation"]
    assert provenance["decision_job_id"] == ids_a["decision"]
    assert ids_b["des"] not in provenance.values()
    assert model["decision"]["status"] == "conditional"


def test_cross_scenario_job_results_are_rejected(db_engine, client):
    analysis_id, dataset_id, scenario_id = _api_workspace(db_engine, "v@example.com")
    login(client, "v@example.com", "pw")
    headers = csrf_header(client)
    _full_chain(client, analysis_id, headers, scenario_id)
    other_id = _add_scenario(db_engine, "v@example.com", analysis_id, dataset_id,
                             "Optimal @ 80%", 0.80)
    with make_sessionmaker(db_engine)() as db:
        user = db.query(User).filter_by(email="v@example.com").one()
        des_job = db.query(Job).filter_by(
            user_id=user.id, kind="workflow_des").order_by(Job.id.desc()).first()
        result = dict(des_job.result_json)
        result["scenario_id"] = other_id
        des_job.result_json = result
        db.commit()
    response = _preview(client, analysis_id, headers)
    assert response.status_code == 409
    assert "scenario" in response.json()["detail"].lower()


def test_malformed_and_missing_evidence_ids_block(db_engine, client):
    analysis_id, _, scenario_id = _api_workspace(db_engine, "w@example.com")
    login(client, "w@example.com", "pw")
    headers = csrf_header(client)
    _full_chain(client, analysis_id, headers, scenario_id)
    with make_sessionmaker(db_engine)() as db:
        user = db.query(User).filter_by(email="w@example.com").one()
        decision = db.query(Job).filter_by(
            user_id=user.id, kind="workflow_decision").order_by(Job.id.desc()).first()
        params = dict(decision.params_json)
        params["validation_job_id"] = "not-an-id"
        decision.params_json = params
        db.commit()
    assert _preview(client, analysis_id, headers).status_code == 409
    with make_sessionmaker(db_engine)() as db:
        user = db.query(User).filter_by(email="w@example.com").one()
        mc = db.query(Job).filter_by(user_id=user.id, kind="workflow_mc").first()
        db.delete(mc)
        db.commit()
    assert _preview(client, analysis_id, headers).status_code == 404


def test_setup_drift_blocks_report(db_engine, client):
    from backend.db.models import AnalysisProject

    analysis_id, _, scenario_id = _api_workspace(db_engine, "x@example.com")
    login(client, "x@example.com", "pw")
    headers = csrf_header(client)
    _full_chain(client, analysis_id, headers, scenario_id)
    with make_sessionmaker(db_engine)() as db:
        analysis = db.get(AnalysisProject, analysis_id)
        setup = dict(analysis.queue_setup_json)
        setup["queue_ids"] = ["lane-A"]
        analysis.queue_setup_json = setup
        db.commit()
    response = _preview(client, analysis_id, headers)
    assert response.status_code == 422


def test_null_uncertainty_and_inactive_lanes_stay_null():
    chain = _chain()
    chain["validation"]["periods"][0]["queues"][0]["mc_failure_rate"] = None
    chain["validation"]["periods"][0]["queues"][0]["validation_verdict"] = "inadequate"
    model = build_separate_report_model(chain)
    queue = model["validation"]["periods"][0]["queues"][0]
    assert queue["mc_failure_rate"] is None
    assert queue["validation_verdict"] == "inadequate"
    assert model["des"]["periods"][0]["lanes"][1]["Wq"] is None
    assert model["cost"]["savings"] is None
    assert model["cost"]["roi"] is None


def test_report_with_forbidden_engines_disabled_still_succeeds(db_engine, client, monkeypatch):
    import inspect

    import backend.api.reports as reports_api
    import backend.api.workflow as workflow_api
    import backend.queueing_engine.services.separate_optimization as sep_opt
    import backend.queueing_engine.simulation.simulation as sim_engine

    def _boom(*args, **kwargs):
        raise AssertionError("analytical engine must not run during reporting")

    analysis_id, _, scenario_id = _api_workspace(db_engine, "y@example.com")
    login(client, "y@example.com", "pw")
    headers = csrf_header(client)
    _full_chain(client, analysis_id, headers, scenario_id)
    monkeypatch.setattr(sep_opt, "optimize_separate", _boom)
    monkeypatch.setattr(sep_opt, "_optimize_separate_des", _boom)
    monkeypatch.setattr(sep_opt, "rank_des_candidates", _boom)
    monkeypatch.setattr(sep_opt, "evaluate_candidate_with_des", _boom)
    monkeypatch.setattr(sim_engine, "mc_simulate_segments", _boom)
    monkeypatch.setattr(sep_opt, "_run_routing_des", _boom)
    monkeypatch.setattr(sim_engine, "simulate_segments_with_trace", _boom)
    monkeypatch.setattr(workflow_api, "validate_selected_plan", _boom)
    monkeypatch.setattr(workflow_api, "_derive_selected_decision", _boom)
    assert _preview(client, analysis_id, headers).status_code == 200
    assert client.get(
        f"/reports/analyses/{analysis_id}/selected/pdf", headers=headers).status_code == 200
    assert client.get(
        f"/reports/analyses/{analysis_id}/selected/excel", headers=headers).status_code == 200
    assert "validate_selected_plan" in inspect.getsource(workflow_api.run_selected_validation)
    assert "optimize_separate" not in inspect.getsource(reports_api)


# --- Execution basis (finding K3) ------------------------------------------------------

def _continuous_chain():
    chain = _chain()
    chain["analysis"] = {**chain["analysis"], "operating_day_hours": 13.0}
    chain["schedule"] = {**chain["schedule"], "utilization_basis": "continuous-day DES",
                         "des": {**chain["schedule"]["des"], "duration_hours": 24.0}}
    return chain


def test_continuous_day_schedule_reports_the_operating_day_not_the_24_hour_default():
    model = build_separate_report_model(_continuous_chain())
    assert model["optimization"]["duration_hours"] == 13.0
    assert "continuous" in model["optimization"]["execution"]
    assert "carry across periods" in model["optimization"]["execution"]
    text = " ".join(model["limitations"])
    assert "continuous run" in text
    assert "queues and breaks carried across period boundaries" in text
    assert "period-independent" not in text


def test_continuous_day_without_a_derivable_day_is_unavailable_not_24():
    chain = _continuous_chain()
    chain["analysis"] = {**chain["analysis"], "operating_day_hours": None}
    assert build_separate_report_model(chain)["optimization"]["duration_hours"] is None


def test_per_date_schedule_keeps_the_period_independent_wording():
    model = build_separate_report_model(_chain())
    assert model["optimization"]["duration_hours"] == 8.0
    assert model["optimization"]["execution"] == "period-independent; no cross-period carryover"
    assert "continuous" not in " ".join(model["limitations"]).lower()


def test_operating_day_hours_spans_the_novamart_periods():
    from backend.api.reports import operating_day_hours
    from backend.data.setup_derivation import derive_setup
    from tests.test_setup_derivation import sheets

    derived = derive_setup(sheets(), None)
    assert derived.errors == []
    labels = sorted({str(row["time"]) for row in derived.records})
    schedule = {"periods": [{"time": label} for label in labels]}
    assert operating_day_hours(derived.setup, schedule) == pytest.approx(13.0)
    assert operating_day_hours(derived.setup, {"periods": [{"time": "no-such-period"}]}) is None
