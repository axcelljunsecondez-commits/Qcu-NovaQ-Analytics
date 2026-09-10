from __future__ import annotations

import openpyxl
import pandas as pd

from backend.reports.report_export import _utilization_status, generate_excel_report, generate_pdf_report


def test_unknown_utilization_is_not_lean():
    assert _utilization_status(None) == 'Unavailable'


def test_pdf_does_not_parse_untrusted_markup(monkeypatch):
    def forbidden_image(*args, **kwargs):
        raise AssertionError('Untrusted markup was parsed as an image')
    monkeypatch.setattr('reportlab.platypus.paraparser.ImageReader', forbidden_image)
    rows = pd.DataFrame([{'time': '<img src="http://audit.invalid/image"/>', 'c_current': 1}])
    assert generate_pdf_report({}, {}, rows, ['<img src="http://audit.invalid/rec"/>']).getvalue().startswith(b'%PDF')


def test_excel_does_not_execute_scenario_text():
    workbook = openpyxl.load_workbook(generate_excel_report(pd.DataFrame([{'time': '=1+1'}])))
    assert workbook['Segments']['A2'].data_type != 'f'


def test_reports_render_incomplete_totals_without_zero_savings():
    rows = pd.DataFrame([{'time': 't', 'c_current': 1, 'c_optimal': None, 'delta_c': None}])
    summary = {'comparison_complete': False, 'total_current_cost': 100,
               'total_optimized_cost': None, 'total_savings': None, 'total_server_change': None}
    workbook = openpyxl.load_workbook(generate_excel_report(rows, summary))
    cells = {workbook['Summary'].cell(r, 1).value: workbook['Summary'].cell(r, 2).value for r in range(1, 25)}
    assert cells['Total Savings'] == 'N/A'
    assert generate_pdf_report({}, summary, rows, []).getvalue().startswith(b'%PDF')
