"""B8 "Formula check" sheet: NovaQ's closed-form queue results as live Excel formulas beside NovaQ's own values.

Spec: docs/superpowers/specs/2026-10-09-b8-reports-pdf-excel.md §4-§5 (owner answers 2026-10-09: OD-UX-12 = A every
closed form live; OD-UX-14 = A helper columns, no SEQUENCE(); OD-B8-4 = A Python re-computation in tests, NOT TESTED
in Excel until the owner's one-time check).

One formula layout serves M/M/1, M/M/c, M/G/c and Parallel M/G/1: the Erlang C waiting time of M/M/c scaled by the
Allen-Cunneen factor (1 + variance * mu^2) / 2 (``queue_models.mgc``). With one counter that product equals the exact
Pollaczek-Khinchine M/G/1 wait, lambda (variance + 1/mu^2) / (2 (1 - rho)); with exponential service the variance is
1/mu^2 and the factor is 1. Checked against ``mm1``, ``mg1``, ``mmc`` and ``mgc`` to 3e-15 relative (2026-10-12).

The Erlang C sum is written with one helper cell per term n = 0 .. c-1 (cells beyond a row's c hold 0), so it works in
every Excel version and in LibreOffice and Google Sheets. Models with finite sums or log-space sums (M/M/c/K, M/G/c/K,
Erlang-A) show NovaQ's values and their steps, marked "formula not available in Excel".
"""

from __future__ import annotations

import math
from collections.abc import Sequence
from typing import Any

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

LIVE_MODELS = frozenset({"M/M/1", "M/M/c", "M/G/c", "Parallel M/G/1"})
EXPONENTIAL_MODELS = frozenset({"M/M/1", "M/M/c"})
# Excel's FACT, like the float factorial in NovaQ's mmc, overflows above 170.
FACT_LIMIT = 170
NOT_AVAILABLE = "formula not available in Excel"
UNSTABLE = "Unstable"
EXCEL_STATUS = (
    "The formulas were checked by re-computing them in Python from the cell contents (OD-B8-4 A). They are NOT TESTED "
    "in Excel until the owner's one-time check is recorded."
)

HEADERS = [
    "Period", "Staffing", "Model", "Customers per hour (lambda)", "Served per counter per hour (mu)", "Counters (c)",
    "Service variance (hours^2)", "Offered load a = lambda/mu", "Busy share rho = a/c", "People waiting Lq (Excel)",
    "Wait in line, minutes (Excel)", "People waiting Lq (NovaQ)", "Wait in line, minutes (NovaQ)",
    "Wait difference, minutes (Excel - NovaQ)", "Note",
]
# Columns of the inputs and results (1-based), then the helper block.
COL = {name: index for index, name in enumerate(
    ["period", "staffing", "model", "lam", "mu", "c", "var", "a", "rho", "lq_x", "wq_x", "lq_n", "wq_n", "diff", "note"],
    start=1)}
HELPER_START = len(HEADERS) + 2  # one empty column between the visible table and the helper cells
FIRST_ROW = 5


def _finite(value: Any) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value) else None


def formula_row(*, period: str, staffing: str, model: str, lam: Any, mu: Any, c: Any, variance: Any,
                novaq_lq: Any, novaq_wq_hours: Any, steps: str | None = None) -> dict[str, Any]:
    """One Formula check row: inputs, NovaQ's values, and whether Excel can reproduce them."""
    lam_value, mu_value, var_value = _finite(lam), _finite(mu), _finite(variance)
    servers = int(c) if isinstance(c, int) and not isinstance(c, bool) else None
    live = (model in LIVE_MODELS and lam_value is not None and lam_value >= 0 and mu_value is not None
            and mu_value > 0 and servers is not None and 1 <= servers <= FACT_LIMIT
            and (model in EXPONENTIAL_MODELS or (var_value is not None and var_value >= 0)))
    wq = _finite(novaq_wq_hours)
    note = None if live else (steps or NOT_AVAILABLE)
    if not live and model in LIVE_MODELS and servers is not None and servers > FACT_LIMIT:
        note = f"{NOT_AVAILABLE}: more than {FACT_LIMIT} counters exceed Excel's FACT"
    return {
        "period": period, "staffing": staffing, "model": model, "lambda": lam_value, "mu": mu_value, "c": servers,
        "variance": None if model in EXPONENTIAL_MODELS else var_value, "live": live,
        "novaq_lq": _finite(novaq_lq), "novaq_wq_minutes": None if wq is None else wq * 60, "note": note,
    }


def helper_count(rows: Sequence[dict[str, Any]]) -> int:
    """Helper cells per row: the largest live counter count (at least one)."""
    return max([row["c"] for row in rows if row["live"]] + [1])


def write_formula_sheet(sheet: Worksheet, rows: Sequence[dict[str, Any]], *, title: str) -> None:
    """Write the Formula check rows; live rows get formulas, the others NovaQ's values and a note."""
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=13)
    sheet["A2"] = EXCEL_STATUS
    sheet["A3"] = ("Live formulas: Erlang C (one helper cell per term, to the right) times the Allen-Cunneen factor "
                   "(1 + variance * mu^2) / 2. Exponential service uses variance = 1/mu^2.")
    terms = helper_count(rows)
    header_row = FIRST_ROW - 1
    for index, header in enumerate(HEADERS, start=1):
        sheet.cell(header_row, index, header).font = Font(bold=True)
    term_cols = [HELPER_START + n for n in range(terms)]
    sum_col, p0_col, lq_mmc_col, factor_col = (HELPER_START + terms + offset for offset in range(4))
    for n, col in enumerate(term_cols):
        sheet.cell(header_row, col, f"Term n={n}: a^n/n!").font = Font(bold=True)
    for col, header in ((sum_col, "Sum of terms"), (p0_col, "P0 (no one present)"),
                        (lq_mmc_col, "Lq with exponential service"), (factor_col, "Allen-Cunneen factor")):
        sheet.cell(header_row, col, header).font = Font(bold=True)

    for offset, row in enumerate(rows):
        r = FIRST_ROW + offset
        cell = {name: f"{get_column_letter(col)}{r}" for name, col in COL.items()}
        values = {"period": row["period"], "staffing": row["staffing"], "model": row["model"], "lam": row["lambda"],
                  "mu": row["mu"], "c": row["c"], "lq_n": row["novaq_lq"],
                  "wq_n": UNSTABLE if row["novaq_wq_minutes"] is None and row["live"] else row["novaq_wq_minutes"],
                  "note": row["note"]}
        for name, value in values.items():
            sheet.cell(r, COL[name], value)
        if not row["live"]:
            sheet.cell(r, COL["var"], row["variance"])
            continue
        lam, mu, c, a, rho = cell["lam"], cell["mu"], cell["c"], cell["a"], cell["rho"]
        variance = f"=1/{mu}^2" if row["model"] in EXPONENTIAL_MODELS else row["variance"]
        sheet.cell(r, COL["var"], variance)
        sheet.cell(r, COL["a"], f"={lam}/{mu}")
        sheet.cell(r, COL["rho"], f"={a}/{c}")
        for n, col in enumerate(term_cols):
            # Excel's 0^0 is #NUM!, so the first term is written as 1.
            term = "1" if n == 0 else f"{a}^{n}/FACT({n})"
            sheet.cell(r, col, f"=IF({n}<{c},{term},0)")
        first, last = f"{get_column_letter(term_cols[0])}{r}", f"{get_column_letter(term_cols[-1])}{r}"
        p0 = f"{get_column_letter(p0_col)}{r}"
        lq_mmc = f"{get_column_letter(lq_mmc_col)}{r}"
        factor = f"{get_column_letter(factor_col)}{r}"
        sheet.cell(r, sum_col, f"=SUM({first}:{last})")
        sheet.cell(r, p0_col, f'=IF({rho}>=1,"{UNSTABLE}",1/({get_column_letter(sum_col)}{r}+{a}^{c}/(FACT({c})*(1-{rho}))))')
        sheet.cell(r, lq_mmc_col, f'=IF({rho}>=1,"{UNSTABLE}",{p0}*{a}^{c}*{rho}/(FACT({c})*(1-{rho})^2))')
        sheet.cell(r, factor_col, f"=(1+{cell['var']}*{mu}^2)/2")
        sheet.cell(r, COL["lq_x"], f'=IF({rho}>=1,"{UNSTABLE}",{lq_mmc}*{factor})')
        sheet.cell(r, COL["wq_x"], f'=IF({rho}>=1,"{UNSTABLE}",IF({lam}=0,0,{cell["lq_x"]}/{lam}*60))')
        sheet.cell(r, COL["diff"], f'=IF(AND(ISNUMBER({cell["wq_x"]}),ISNUMBER({cell["wq_n"]})),{cell["wq_x"]}-{cell["wq_n"]},"")')
