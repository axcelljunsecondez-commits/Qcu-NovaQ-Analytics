"""Presentation regressions on real isolated report routes; all data is SYNTHETIC."""
from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook

from backend.reports.plan_report import _format_report_workbook
from tests import test_b8_full_report as full
from tests import test_report_evidence_fixes as repairs


@pytest.mark.parametrize('kind,level', [('shared','plan'), ('shared','full'),
                                         ('full_day','full'), ('separate','plan'), ('separate','full')])
def test_report_exports_have_readable_detail_and_wrapped_descriptions(client, db_engine, kind, level):
    if kind == 'full_day':
        ids, _, _ = repairs._full_day_chain(client, db_engine)
        analysis_id = ids['analysis_id']
    elif kind == 'separate':
        analysis_id = full._separate_chain(client, db_engine)
    else:
        analysis_id = full._shared_chain(client, db_engine)
    response = client.get(f'/analyses/{analysis_id}/reports/{level}.xlsx')
    assert response.status_code == 200, response.text
    book = load_workbook(io.BytesIO(response.content))
    assert book['Overview'].column_dimensions['A'].width >= 35
    assert book['Overview']['B11'].alignment.wrap_text is True
    assert book['Your data'].column_dimensions['A'].width >= 25
    for name in ('Plan', 'Plan search'):
        sheet = book[name]
        assert sheet['A3'].alignment.wrap_text is True
        assert sheet.freeze_panes is not None
    formula = book['Formula check']
    assert formula['E4'].alignment.wrap_text is True
    assert formula.row_dimensions[4].height >= 48
    assert formula['N6'].number_format == '0.00E+00'
    assert formula.freeze_panes == 'D5'
    assert book['Plan']['D4'].number_format == '0.0%'
    assert book['Plan']['F5'].number_format == '0.0000'
    assert book['Plan']['I4'].alignment.indent == 1
    assert book['Overview']['B10'].number_format == '#,##0.00'
    assert book['Assumptions'].column_dimensions['A'].width >= 35
    assert book['Provenance']['B7'].alignment.wrap_text is True
    if level == 'full':
        assert book['Decision']['B6'].alignment.wrap_text is True
        assert book['Decision'].row_dimensions[6].height >= 32
        assert book['Checks']['B4'].alignment.wrap_text is True
        simulation = book['Simulation']
        for row in range(4, simulation.max_row):
            if simulation.cell(row, 2).value == 'percent':
                assert simulation.cell(row + 1, 2).number_format == '0.0"%"'


def test_formatting_preserves_formulas_zero_missing_and_populated_title_neighbors():
    book = Workbook()
    sheet = book.active
    sheet.title = 'Overview'
    for row in [['Title', 0], [], ['Item', 'Value'], ['Missing', None], ['Zero', 0],
                ['Unstable', 'Unstable'], ['Formula', '=IF(B5=0,0,B5/$B$1)']]:
        sheet.append(row)
    def payload():
        return [(s.title, [(c.coordinate, c.data_type, c.value) for row in s for c in row if c.value is not None])
                for s in book]
    before = payload()
    _format_report_workbook(book)
    assert payload() == before
    assert sheet['B4'].value is None
    assert sheet['B5'].value == 0
    assert sheet['B1'].value == 0
