"""B8 step 15b: the plan and full report PDFs and the full report (level 2).

Spec: docs/superpowers/specs/2026-10-09-b8-reports-pdf-excel.md §1, §6 tests 1-3 (owner answers 2026-10-09).
Both line types go through the public routes: plan, selection, simulation, checks and Decision.

Data: SYNTHETIC (rates and setups made for the test).
"""

from __future__ import annotations

import base64
import io
import re
import statistics
import zlib
from typing import Any

from openpyxl import load_workbook

from backend.db.models import Dataset
from backend.reports.full_report import SHARED_SOURCE
from backend.reports.report_pdf import generate_report_pdf, pdf_sections
from tests import test_b13_per_date_day as b13
from tests.helpers import create_user, csrf_header, login, make_sessionmaker

SHARED_SETUP = {"queue_structure": "shared_queue", "fixed_server_count": 3, "staffing_varies_by_period": False,
                "capacity_mode": "unlimited", "abandonment_mode": "not_modeled"}
CSV = b"time,lambda,mu,c\n08:00-09:00,40,15,3\n09:00-10:00,20,15,3\n"
OPTIONS = {"target_utilization": 0.85, "server_cost_per_hr": 87.0, "customer_waiting_cost": 100.0, "max_servers": 6}


def pdf_text(content: bytes) -> str:
    """The text shown by a ReportLab PDF: every page stream decompressed, the (strings) of its text operators joined."""
    assert content.startswith(b"%PDF")
    parts = []
    # ReportLab writes each stream as ASCII85 (ending "~>") of Flate data, with no newline before "endstream".
    for raw in re.findall(rb"stream\r?\n(.*?)endstream", content, re.S):
        data = raw.strip()
        try:
            stream = zlib.decompress(base64.a85decode(data[:-2] if data.endswith(b"~>") else data))
        except (ValueError, zlib.error):
            continue
        parts += [item.decode("latin-1") for item in re.findall(rb"\((.*?)(?<!\\)\)\s*Tj", stream, re.S)]
    return " ".join(parts).replace("\\(", "(").replace("\\)", ")")


def _shared_chain(client, db_engine, *, decide: bool = True) -> int:
    create_user(db_engine, "owner@example.com", "pw")
    assert login(client, "owner@example.com", "pw") == 200
    headers = csrf_header(client)
    analysis = client.post("/analyses", headers=headers, json={"name": "Counters", "queue_setup": SHARED_SETUP}).json()["analysis"]
    a = analysis["id"]
    dataset = client.post(f"/analyses/{a}/datasets", headers=headers, files={"file": ("r.csv", CSV, "text/csv")}).json()["dataset"]
    segments = [{key: row[key] for key in ("time", "lambda", "mu", "c")} for row in dataset["normalized"]]
    results = client.post("/optimize/batch", headers=headers, json={"segments": segments, **OPTIONS}).json()
    scenario = client.post("/scenarios", headers=headers, json={
        "name": "Four at eight", "analysis_id": a, "dataset_id": dataset["id"], "results": results,
        "settings": {**OPTIONS, "calculation": {"schema_version": 1, "engine_version": "novaq-2026-09-system-v2",
                                                "input_segments": segments, "options": OPTIONS, "what_if_multiplier": 1,
                                                "calculated_at": "2026-10-13T00:00:00+00:00"}}}).json()["scenario"]
    assert client.post(f"/analyses/{a}/workflow/selection", headers=headers, json={"scenario_id": scenario["id"]}).status_code == 200
    assert client.post(f"/analyses/{a}/workflow/simulation/des", headers=headers, json={"sim_hours": 8, "seed": 7}).status_code == 200
    assert client.post(f"/analyses/{a}/workflow/simulation/mc", headers=headers, json={"num_trials": 200, "seed": 7}).status_code == 200
    assert client.post(f"/analyses/{a}/workflow/simulation/validation", headers=headers,
                       json={"des_sim_hours": 8, "mc_trials": 200}).status_code == 200
    if decide:
        assert client.post(f"/analyses/{a}/workflow/decision", headers=headers).status_code == 200
    return a


def _model(client, analysis_id: int, level: str = "full") -> Any:
    return client.get(f"/analyses/{analysis_id}/reports/{level}/preview")


def test_the_plan_pdf_is_plain_words_without_decision_words(db_engine, client):
    a = _shared_chain(client, db_engine, decide=False)
    response = client.get(f"/analyses/{a}/reports/plan.pdf")
    assert response.status_code == 200 and response.headers["content-type"] == "application/pdf"
    text = pdf_text(response.content)
    for phrase in ("NovaQ plan report", "Calculated result", "not yet checked by simulation", "Key results: now and with the plan",
                   "Counters now and with the plan", "Four at eight", "08:00-09:00: 3 to 4"):
        assert phrase in text, phrase
    assert "Decision" not in text and "Wq" not in text and "rho" not in text.lower()


def test_the_shared_full_report_needs_a_current_decision(db_engine, client):
    a = _shared_chain(client, db_engine, decide=False)
    refused = _model(client, a)
    assert refused.status_code == 422 and refused.json()["detail"]["code"] == "no_current_decision"
    assert client.get(f"/analyses/{a}/reports/full.pdf").status_code == 422
    assert client.post(f"/analyses/{a}/workflow/decision", headers=csrf_header(client)).status_code == 200
    assert _model(client, a).status_code == 200
    # A new simulation run makes the Decision stale, so the full report is refused again.
    client.post(f"/analyses/{a}/workflow/simulation/des", headers=csrf_header(client), json={"sim_hours": 8, "seed": 8})
    assert _model(client, a).status_code == 422


def test_the_shared_full_report_quotes_the_stored_evidence(db_engine, client):
    a = _shared_chain(client, db_engine)
    workflow = client.get(f"/analyses/{a}/workflow").json()
    decision = workflow["decision"]["result"]
    model = _model(client, a).json()["model"]
    assert model["report"] == "full" and model["cover"]["level"] == "full"
    assert {key: model["decision"][key] for key in ("status", "headline", "recommendation", "rationale")} == {
        key: decision[key] for key in ("status", "headline", "recommendation", "rationale")}
    simulation = model["simulation"]
    assert simulation["source"] == SHARED_SOURCE
    summary = workflow["des"]["result"]["summary"]
    assert simulation["headlines"][0]["value"] == summary["average_wait_hours"] * 60
    assert simulation["headlines"][2]["value"] is None  # the period-by-period run does not count served customers
    mc_rows = workflow["mc"]["result"]["results"]
    assert [p["value"] for p in simulation["charts"][1]["points"]] == [row["failure_rate"] * 100 for row in mc_rows]
    checks = {item["check"]: item for item in model["checks"]}
    assert checks["Many variations"]["settings"].startswith("200 trials")
    assert checks["Simulation run"]["settings"] == "8.0 simulated hours per period, seed 7"
    passed = sum(1 for row in mc_rows if row["status"] == "PASS")
    assert checks["Many variations"]["outcome"] == f"{passed} of {len(mc_rows)} within the allowed failure share"
    assert any("full-day run" in line for line in model["not_included"])

    pdf = pdf_text(client.get(f"/analyses/{a}/reports/full.pdf").content)
    for phrase in ("NovaQ full report", "Decision", "Simulation key results", "Checks", "What is not included",
                   "not recorded by this run"):
        assert phrase in pdf, phrase
    workbook = load_workbook(io.BytesIO(client.get(f"/analyses/{a}/reports/full.xlsx").content))
    assert workbook.sheetnames == ["Overview", "Your data", "Plan", "Formula check", "Plan search", "Simulation", "Checks",
                                   "Decision", "Assumptions", "Provenance"]
    assert workbook["Decision"]["B5"].value == decision["headline"]


def _separate_chain(client, db_engine, *, decide: bool = True) -> int:
    analysis_id, dataset_id, rows = b13._workspace(db_engine, "representative_day")
    with make_sessionmaker(db_engine)() as db:
        dataset = db.get(Dataset, dataset_id)
        dataset.normalized_json = [{**row, "variance": statistics.pvariance(row["service_samples_hours"])}
                                   for row in dataset.normalized_json]
        db.commit()
    login(client, "owner@example.com", "pw")
    out = b13._optimize(client, analysis_id)
    saved = b13._save(client, analysis_id, dataset_id, b13._settings(analysis_id, dataset_id, rows, out["engine_version"]),
                      out["schedule"]).json()["scenario"]
    h = csrf_header(client)
    assert client.post(f"/analyses/{analysis_id}/workflow/selection", headers=h, json={"scenario_id": saved["id"]}).status_code == 200
    assert client.post(f"/analyses/{analysis_id}/workflow/simulation/des/selected", headers=h,
                       json={"seed": 7, "load_replications": 2}).status_code == 200
    assert client.post(f"/analyses/{analysis_id}/workflow/simulation/mc/selected", headers=h,
                       json={"num_trials": 200, "seed": 7}).status_code == 200
    assert client.post(f"/analyses/{analysis_id}/workflow/simulation/validation/selected", headers=h, json={}).status_code == 200
    if decide:
        assert client.post(f"/analyses/{analysis_id}/workflow/decision/selected", headers=h, json={}).status_code == 200
    return analysis_id


def test_the_separate_full_report_keeps_todays_chain_and_quotes_it(db_engine, client):
    a = _separate_chain(client, db_engine, decide=False)
    refused = _model(client, a)
    assert refused.status_code == 404 and "Decision" in str(refused.json()["detail"])  # today's chain reason, unchanged
    assert client.post(f"/analyses/{a}/workflow/decision/selected", headers=csrf_header(client), json={}).status_code == 200
    model = _model(client, a).json()["model"]
    assert model["line_type"] == "separate"
    checks = {item["check"]: item for item in model["checks"]}
    assert checks["Results check"]["outcome"] == "pass: 4 of 4 line-periods passed, 0 failed, 0 without enough data"
    assert checks["Many variations"]["outcome"] == "4 of 4 within the allowed failure share"
    assert "whole day as one continuous run" in model["simulation"]["source"]
    assert model["simulation"]["headlines"][2]["value"] == 6  # customers served on the replayed day
    assert model["decision"]["status"] == "conditional"
    assert client.get(f"/analyses/{a}/reports/full.pdf").status_code == 200
    assert "NovaQ plan report" in pdf_text(client.get(f"/analyses/{a}/reports/plan.pdf").content)


def test_other_users_cannot_read_any_report_file(db_engine, client):
    a = _shared_chain(client, db_engine)
    client.post("/auth/logout", headers=csrf_header(client))
    create_user(db_engine, "other@example.com", "pw")
    login(client, "other@example.com", "pw")
    for path in ("plan.pdf", "full/preview", "full.pdf", "full.xlsx"):
        assert client.get(f"/analyses/{a}/reports/{path}").status_code == 404, path


def test_the_listed_contents_are_exactly_the_files_contents(db_engine, client):
    """Spec §6 test 1: the page lists what the files contain, from the same rules."""
    a = _shared_chain(client, db_engine)
    for level in ("plan", "full"):
        preview = _model(client, a, level).json()
        contents = preview["contents"]
        text = pdf_text(client.get(f"/analyses/{a}/reports/{level}.pdf").content)
        assert all(section in text for section in contents["pdf"]), level
        workbook = load_workbook(io.BytesIO(client.get(f"/analyses/{a}/reports/{level}.xlsx").content))
        assert workbook.sheetnames == contents["excel"], level
    assert "Decision" not in _model(client, a, "plan").json()["contents"]["pdf"]
    plan_preview, full_preview = _model(client, a, "plan").json(), _model(client, a, "full").json()
    for kind in ("pdf", "excel"):
        # The full report has the plan report's parts plus the ones marked "after the checks".
        assert sorted(plan_preview["contents"][kind] + plan_preview["after_checks"][kind]) == sorted(full_preview["contents"][kind])
    assert full_preview["after_checks"] == {"pdf": [], "excel": []}
    # A Setup with breaks adds its section; generate_report_pdf asserts its headings equal pdf_sections.
    with_breaks = {**full_preview["model"], "breaks": [{"scheduled_start_time": "08:30", "duration_minutes": 15}]}
    assert pdf_sections(with_breaks) == [*full_preview["contents"]["pdf"][:2], "Breaks in the Setup", *full_preview["contents"]["pdf"][2:]]
    assert "Breaks in the Setup" in pdf_text(generate_report_pdf(with_breaks).getvalue())


def test_previews_have_their_own_limit_so_browsing_reports_does_not_block_downloads(db_engine, client):
    """U10 F14: the Reports page asks for two previews per visit; they no longer use up the file limit."""
    a = _shared_chain(client, db_engine)
    settings = client.app.state.settings
    settings.rate_limit_report, settings.rate_limit_report_preview = 1, 4
    assert client.get(f"/analyses/{a}/reports/full.pdf").status_code == 200
    assert client.get(f"/analyses/{a}/reports/full.xlsx").status_code == 429  # one file per window
    for level in ("plan", "full", "plan", "full"):
        assert _model(client, a, level).status_code == 200, level
    refused = _model(client, a, "plan")
    assert refused.status_code == 429 and "Retry-After" in refused.headers  # the preview limit is still enforced
