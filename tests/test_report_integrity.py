from __future__ import annotations

import openpyxl
import pandas as pd

from backend.reports.report_export import _utilization_status, generate_excel_report, generate_pdf_report


def test_unknown_utilization_is_not_lean():
    assert _utilization_status(None) == 'Unavailable'


def test_pdf_does_not_parse_untrusted_markup(monkeypatch):
    def forbidden_image(*args, **kwargs):
        raise AssertionError('Untrusted markup was parsed as an image')
    monkeypatch.setattr('reportlab.platypus.paraparser.ImageReader', forbidden_image)
    rows = pd.DataFrame([{'time': '<img src="http://audit.invalid/image"/>', 'c_current': 1}])
    assert generate_pdf_report({}, {}, rows, ['<img src="http://audit.invalid/rec"/>']).getvalue().startswith(b'%PDF')


def test_excel_does_not_execute_scenario_text():
    workbook = openpyxl.load_workbook(generate_excel_report(pd.DataFrame([{'time': '=1+1'}])))
    assert workbook['Segments']['A2'].data_type != 'f'


def test_reports_render_incomplete_totals_without_zero_savings():
    rows = pd.DataFrame([{'time': 't', 'c_current': 1, 'c_optimal': None, 'delta_c': None}])
    summary = {'comparison_complete': False, 'total_current_cost': 100,
               'total_optimized_cost': None, 'total_savings': None, 'total_server_change': None}
    workbook = openpyxl.load_workbook(generate_excel_report(rows, summary))
    cells = {workbook['Summary'].cell(r, 1).value: workbook['Summary'].cell(r, 2).value for r in range(1, 25)}
    assert cells['Total Savings'] == 'N/A'
    assert generate_pdf_report({}, summary, rows, []).getvalue().startswith(b'%PDF')


def _blocked_rows() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                'time': '08:00', 'c_current': 1, 'c_optimal': None,
                'rho_current': None, 'rho_optimal': None,
                'Wq_current': None, 'Wq_optimal': None, 'delta_c': None,
            },
            {
                'time': '09:00', 'c_current': None, 'c_optimal': None,
                'rho_current': None, 'rho_optimal': None,
                'Wq_current': None, 'Wq_optimal': None, 'delta_c': None,
            },
        ]
    )


def _recorded_paragraph_texts(monkeypatch, report_module) -> list:
    texts: list = []
    original_paragraph = report_module.Paragraph

    def recording_paragraph(text, *args, **kwargs):
        texts.append(text)
        return original_paragraph(text, *args, **kwargs)

    monkeypatch.setattr(report_module, 'Paragraph', recording_paragraph)
    return texts


def test_pdf_staffing_cells_use_na_for_missing_values(monkeypatch):
    import backend.reports.report_export as report_export

    texts = _recorded_paragraph_texts(monkeypatch, report_export)
    report_export.generate_pdf_report({}, {'comparison_complete': False}, _blocked_rows(), [])
    assert 'None' not in texts
    assert 'nan' not in texts
    assert texts.count('N/A') >= 2
    assert '1' in texts


def test_pdf_staffing_cells_render_integral_floats_exactly(monkeypatch):
    import backend.reports.report_export as report_export

    texts = _recorded_paragraph_texts(monkeypatch, report_export)
    rows = pd.DataFrame(
        [{'time': '08:00', 'c_current': 3.0, 'c_optimal': 4.0,
          'rho_current': None, 'rho_optimal': None,
          'Wq_current': None, 'Wq_optimal': None, 'delta_c': 1.0}]
    )
    report_export.generate_pdf_report({}, {'comparison_complete': False}, rows, [])
    assert '3' in texts
    assert '4' in texts
    assert '3.0' not in texts
    assert '4.0' not in texts


def test_excel_total_server_change_na_when_unavailable():
    summary = {'comparison_complete': False, 'total_current_cost': None,
               'total_optimized_cost': None, 'total_savings': None, 'total_server_change': None}
    workbook = openpyxl.load_workbook(generate_excel_report(_blocked_rows(), summary))
    cells = {workbook['Summary'].cell(r, 1).value: workbook['Summary'].cell(r, 2).value for r in range(1, 25)}
    assert cells['Total Server Change'] == 'N/A'


def test_excel_total_server_change_zero_when_genuine():
    summary = {'comparison_complete': True, 'total_current_cost': 100,
               'total_optimized_cost': 100, 'total_savings': 0, 'total_server_change': 0}
    workbook = openpyxl.load_workbook(generate_excel_report(_blocked_rows(), summary))
    cells = {workbook['Summary'].cell(r, 1).value: workbook['Summary'].cell(r, 2).value for r in range(1, 25)}
    assert cells['Total Server Change'] == '0'
