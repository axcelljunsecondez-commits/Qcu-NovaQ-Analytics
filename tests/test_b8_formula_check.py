"""B8 Formula check sheet: live Excel formulas re-computed in Python equal NovaQ's own values.

Spec: docs/superpowers/specs/2026-10-09-b8-reports-pdf-excel.md §6 tests 4 and 5 (OD-UX-12 A, OD-UX-14 A, OD-B8-4 A).
Data: SYNTHETIC reference cases (stable, near-saturated, one counter, many counters, no arrivals, unstable).
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from backend.queueing_engine.models.queue_models import mg1, mgc, mm1, mmc, mmck
from backend.reports.formula_check import COL, FIRST_ROW, NOT_AVAILABLE, UNSTABLE, formula_row, write_formula_sheet
from tests.excel_formula_eval import FormulaError, SheetEvaluator

CASES = [  # (model, lambda, mu, c, variance, NovaQ result)
    ("M/M/1", 10.0, 15.0, 1, None, mm1(10.0, 15.0)),
    ("M/M/c", 40.0, 15.0, 3, None, mmc(40.0, 15.0, 3)),
    ("M/M/c", 44.9, 15.0, 3, None, mmc(44.9, 15.0, 3)),  # near saturation, rho = 0.99778
    ("M/M/c", 1400.0, 15.0, 100, None, mmc(1400.0, 15.0, 100)),
    ("M/G/c", 40.0, 15.0, 3, 0.01, mgc(40.0, 15.0, 3, 0.01)),
    ("M/G/c", 40.0, 15.0, 3, 0.0, mgc(40.0, 15.0, 3, 0.0)),  # deterministic service
    ("Parallel M/G/1", 10.0, 15.0, 1, 0.002, mg1(10.0, 15.0, 0.002)),
    ("M/M/c", 0.0, 15.0, 2, None, mmc(0.0, 15.0, 2)),  # no arrivals
]


def _sheet(rows):
    workbook = Workbook()
    sheet = workbook.active
    write_formula_sheet(sheet, rows, title="Formula check")
    return sheet, SheetEvaluator(sheet)


def _rows():
    return [formula_row(period=f"P{index}", staffing="Plan", model=model, lam=lam, mu=mu, c=c, variance=variance,
                        novaq_lq=result["Lq"], novaq_wq_hours=result["Wq"])
            for index, (model, lam, mu, c, variance, result) in enumerate(CASES)]


@pytest.mark.parametrize("index", range(len(CASES)))
def test_each_live_formula_recomputes_novaqs_value(index):
    sheet, evaluator = _sheet(_rows())
    r = FIRST_ROW + index
    model, _, _, _, _, result = CASES[index]
    wq_excel = evaluator.value(f"{get_column_letter(COL['wq_x'])}{r}")
    lq_excel = evaluator.value(f"{get_column_letter(COL['lq_x'])}{r}")
    assert wq_excel == pytest.approx(result["Wq"] * 60, rel=1e-9, abs=1e-12), model
    assert lq_excel == pytest.approx(result["Lq"], rel=1e-9, abs=1e-12), model
    assert evaluator.value(f"{get_column_letter(COL['rho'])}{r}") == pytest.approx(result["rho"], rel=1e-12)
    difference = evaluator.value(f"{get_column_letter(COL['diff'])}{r}")
    assert abs(difference) <= 1e-9 * max(1.0, result["Wq"] * 60)
    # The cells really are formulas, not pasted values.
    assert str(sheet[f"{get_column_letter(COL['wq_x'])}{r}"].value).startswith("=")


def test_an_unstable_period_says_unstable_and_never_shows_a_number():
    result = mmc(50.0, 15.0, 3)
    rows = [formula_row(period="P", staffing="Now", model="M/M/c", lam=50.0, mu=15.0, c=3, variance=None,
                        novaq_lq=result["Lq"], novaq_wq_hours=result["Wq"])]
    _, evaluator = _sheet(rows)
    for key in ("lq_x", "wq_x", "wq_n"):
        assert evaluator.value(f"{get_column_letter(COL[key])}{FIRST_ROW}") == UNSTABLE
    assert evaluator.value(f"{get_column_letter(COL['diff'])}{FIRST_ROW}") == ""


def test_models_without_a_closed_form_show_novaqs_values_and_say_why():
    result = mmck(40.0, 15.0, 2, 6)
    row = formula_row(period="P", staffing="Plan", model="M/M/c/K", lam=40.0, mu=15.0, c=2, variance=None,
                      novaq_lq=result["Lq"], novaq_wq_hours=result["Wq"],
                      steps=f"{NOT_AVAILABLE}; turned away {result['blocking_probability']:.4f}")
    big = formula_row(period="Q", staffing="Plan", model="M/M/c", lam=1.0, mu=1.0, c=200, variance=None,
                      novaq_lq=None, novaq_wq_hours=None)
    sheet, evaluator = _sheet([row, big])
    assert row["live"] is False and big["live"] is False
    assert sheet.cell(FIRST_ROW, COL["note"]).value.startswith(NOT_AVAILABLE)
    assert "more than 170 counters" in sheet.cell(FIRST_ROW + 1, COL["note"]).value
    assert evaluator.value(f"{get_column_letter(COL['wq_n'])}{FIRST_ROW}") == pytest.approx(result["Wq"] * 60)
    assert sheet.cell(FIRST_ROW, COL["wq_x"]).value is None  # no formula pretends to compute it


def test_helper_cells_follow_each_rows_counters_and_never_use_sequence():
    sheet, evaluator = _sheet(_rows())
    formulas = [cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert formulas and not any("SEQUENCE" in formula for formula in formulas)
    # Row of the c = 3 case: terms n = 0, 1, 2 are a^n/n!, n >= 3 are 0.
    from backend.reports.formula_check import HELPER_START
    r = FIRST_ROW + 1
    a = 40.0 / 15.0
    terms = [evaluator.value(f"{get_column_letter(HELPER_START + n)}{r}") for n in range(5)]
    assert terms == pytest.approx([1.0, a, a * a / 2, 0.0, 0.0])


def test_the_evaluator_refuses_what_it_does_not_support():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "=SEQUENCE(3)"
    sheet["A2"] = "=0^0"
    with pytest.raises(FormulaError):
        SheetEvaluator(sheet).value("A1")
    with pytest.raises(FormulaError):
        SheetEvaluator(sheet).value("A2")
