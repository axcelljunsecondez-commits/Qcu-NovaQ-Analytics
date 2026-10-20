"""Owner-authorized report repairs. All records and rates here are SYNTHETIC.

Exercise isolated API workflows and actual file renderers; never contact the running
Docker database. Preserve stored calculations, Decision gates and numerical values.
"""

from __future__ import annotations

import io
from copy import deepcopy

import pytest
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch

from backend.api import plan_reports
from backend.db.models import AnalysisProject, Job, Scenario
from backend.reports.full_report import shared_full_sections
from backend.reports.report_pdf import _styles, _table
from tests import test_b8_full_report as full
from tests import test_b8_plan_report as plan
from tests import test_od_ux_13_full_day_decision as day
from tests.helpers import csrf_header, make_sessionmaker


def _full_day_chain(client, db_engine):
    ids = plan._shared_plan(client, db_engine)
    a = ids['analysis_id']
    assert client.put(f'/analyses/{a}/closing-policy', headers=csrf_header(client),
                      json={'closing_policy': 'DRAIN'}).status_code == 200
    plan._select(client, ids)
    run = day._day_run(client, a)['evidence']
    day._variations(client, a)
    decision = day._decide(client, a)
    assert decision['basis'] == 'full_day_run'
    return ids, run, decision


def test_full_day_files_quote_the_matching_run_without_legacy_evidence(client, db_engine):
    ids, run, decision = _full_day_chain(client, db_engine)
    a = ids['analysis_id']
    workflow = day._workflow(client, a)
    assert workflow['des'] is None and workflow['validation'] is None
    with make_sessionmaker(db_engine)() as db:
        before_run = deepcopy(db.get(Job, run['id']).result_json)
        before_plan = deepcopy(db.get(Scenario, ids['scenario_id']).results_json)
    response = full._model(client, a)
    assert response.status_code == 200, response.text
    model = response.json()['model']
    simulation = model['simulation']
    assert simulation['basis'] == 'full_day_run'
    assert model['decision']['recommendation'] == decision['recommendation']
    summary = run['result']['summary']
    expected = [summary['waiting_time']['customer_weighted_mean_hours']['value'] * 60,
                summary['metrics']['max_queue_in_horizon']['mean'],
                summary['customers']['per_replication']['served']['mean'],
                summary['customers']['per_replication']['unserved']['mean']]
    assert [item['value'] for item in simulation['headlines']] == pytest.approx(expected)
    assert 'served customers' in simulation['headlines'][0]['label']
    assert all('average per day' in item['label'] for item in simulation['headlines'][1:])
    chart = simulation['charts'][0]
    assert chart['unit'] == 'percent' and 'busy' in chart['title'].lower()
    assert [p['value'] for p in chart['points']] == pytest.approx([
        s['summary']['mean'] * 100 for s in run['result']['period_utilization']['segments']])
    checks = {c['check']: c for c in model['checks']}
    assert checks['Simulation run']['outcome'] == 'done'
    assert f"run {run['id']}" in checks['Simulation run']['settings']
    assert '5 repeated days' in checks['Simulation run']['settings']
    assert 'seed 5' in checks['Simulation run']['settings']
    assert 'DRAIN' in checks['Simulation run']['settings']
    assert 'period-by-period validation' in checks['Results check']['settings']
    text = full.pdf_text(client.get(f'/analyses/{a}/reports/full.pdf').content)
    assert 'Full-day simulation' in text
    assert 'preview' not in text and 'no pass or fail rule' not in text
    assert 'not recorded by this run' not in text and 'None simulated hours' not in text
    workbook = load_workbook(io.BytesIO(client.get(f'/analyses/{a}/reports/full.xlsx').content))
    assert [workbook['Simulation'].cell(r, 2).value for r in range(4, 8)] == pytest.approx(expected)
    assert workbook['Decision']['B5'].value == decision['headline']
    assert workbook['Simulation'].column_dimensions['A'].width >= 60
    assert workbook['Checks']['B4'].alignment.wrap_text
    with make_sessionmaker(db_engine)() as db:
        assert db.get(Job, run['id']).result_json == before_run
        assert db.get(Scenario, ids['scenario_id']).results_json == before_plan


@pytest.mark.parametrize('change', ['new_run', 'missing_run', 'setup', 'dataset'])
def test_full_day_reports_keep_the_existing_freshness_gate(client, db_engine, change):
    ids, run, _ = _full_day_chain(client, db_engine)
    a = ids['analysis_id']
    if change == 'new_run':
        day._day_run(client, a, seed=9)
    elif change == 'dataset':
        response = client.post(f'/analyses/{a}/datasets', headers=csrf_header(client),
                               files={'file': ('replacement.csv', plan.CSV, 'text/csv')})
        assert response.status_code == 201, response.text
    else:
        with make_sessionmaker(db_engine)() as db:
            if change == 'missing_run':
                db.delete(db.get(Job, run['id']))
            else:
                analysis = db.get(AnalysisProject, a)
                analysis.queue_setup_json = {**analysis.queue_setup_json, 'fixed_server_count': 3}
            db.commit()
    for path in ('full/preview', 'full.pdf', 'full.xlsx'):
        response = client.get(f'/analyses/{a}/reports/{path}')
        assert response.status_code == 422, (change, path, response.text)


def test_reports_refuse_when_evidence_changes_between_the_gate_and_render_read(client, db_engine, monkeypatch):
    ids, _, _ = _full_day_chain(client, db_engine)
    read = plan_reports._current_evidence

    def invalidated(*args):
        evidence = read(*args)
        evidence['decision'] = None  # simulate invalidation at the second read boundary
        return evidence

    monkeypatch.setattr(plan_reports, '_current_evidence', invalidated)
    for path in ('full/preview', 'full.pdf', 'full.xlsx'):
        response = client.get(f"/analyses/{ids['analysis_id']}/reports/{path}")
        assert response.status_code == 422
        assert response.json()['detail']['code'] == 'no_current_decision'


def test_period_report_keeps_its_basis_without_claiming_all_full_days_are_previews(client, db_engine):
    a = full._shared_chain(client, db_engine)
    model = full._model(client, a).json()['model']
    assert model['simulation']['source'].startswith('Period-by-period simulation')
    text = full.pdf_text(client.get(f'/analyses/{a}/reports/full.pdf').content)
    assert 'no pass or fail rule' not in text and 'is a preview' not in text
    assert model['simulation']['headlines'][2]['value'] is None


def test_what_if_workbook_identifies_saved_calculation_inputs_without_changing_them(client, db_engine):
    ids = plan._shared_plan(client, db_engine)
    segments = [{**row, 'lambda': row['lambda'] * 1.2} for row in ids['segments']]
    results = client.post('/optimize/batch', headers=csrf_header(client),
                          json={'segments': segments, **plan.OPTIONS}).json()
    saved = client.post('/scenarios', headers=csrf_header(client), json={
        'name': 'SYNTHETIC demand 1.2', 'analysis_id': ids['analysis_id'], 'dataset_id': ids['dataset_id'],
        'results': results, 'settings': {**plan.OPTIONS, 'calculation': {
            'schema_version': 1, 'engine_version': 'novaq-2026-09-system-v2', 'input_segments': segments,
            'options': plan.OPTIONS, 'what_if_multiplier': 1.2, 'calculated_at': '2026-10-16T00:00:00+00:00'}}})
    assert saved.status_code == 201, saved.text
    scenario = saved.json()['scenario']
    plan._select(client, {**ids, 'scenario_id': scenario['id']})
    response = client.get(f"/analyses/{ids['analysis_id']}/reports/plan.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook['Your data']['A1'].value == 'Saved calculation inputs'
    assert [workbook['Your data'][f'B{r}'].value for r in (4, 5)] == [48, 24]
    assert dict(list(workbook['Assumptions'].values)[3:])['what_if_multiplier'] == 1.2
    with make_sessionmaker(db_engine)() as db:
        assert db.get(Scenario, scenario['id']).settings_json['calculation']['input_segments'] == segments


def test_separate_comparison_identifies_optimizer_basis_and_keeps_values(client, db_engine):
    a = full._separate_chain(client, db_engine)
    model = full._model(client, a).json()['model']
    assert model['decision']['status'] == 'conditional'
    assert 'Current' in model['decision']['recommendation']
    assert 'optimizer simulation' in model['key_results']['cost_basis'].lower()
    assert 'savings are unavailable' in model['key_results']['cost_basis'].lower()
    text = full.pdf_text(client.get(f'/analyses/{a}/reports/full.pdf').content)
    assert 'Current staffing (simulated)' in text
    assert 'no comparable Current cost basis' in text
    workbook = load_workbook(io.BytesIO(client.get(f'/analyses/{a}/reports/full.xlsx').content))
    overview = dict(list(workbook['Overview'].values)[3:])
    # openpyxl serializes numeric cells with %.16g; allow only that text round-trip rounding.
    assert overview['Cost at current staffing (optimizer simulation)'] == pytest.approx(model['key_results']['cost_now'], rel=1e-14, abs=1e-12)
    assert 'simulated' in workbook['Plan']['H3'].value.lower()
    assert workbook['Plan']['H3'].alignment.wrap_text
    assert workbook['Plan'].column_dimensions['H'].width >= 24
    assert workbook['Overview']['B11'].alignment.wrap_text
    for r, period in enumerate(model['periods'], start=4):
        assert workbook['Plan'].cell(r, 8).value == pytest.approx(period['cost_now'], rel=1e-14, abs=1e-12)
        assert workbook['Plan'].cell(r, 9).value == pytest.approx(period['cost_plan'], rel=1e-14, abs=1e-12)
    assert workbook['Your data']['A1'].value == 'Your data (the planning file)'
    assert 'Decision' not in full.pdf_text(client.get(f'/analyses/{a}/reports/plan.pdf').content)


def test_missing_full_day_statistics_remain_unavailable_and_insufficient_is_not_passed():
    decision = {'basis': 'full_day_run', 'evidence_ids': {'continuous': 12},
                'facts': {'periods': [{'time': '08:00-09:00', 'verdict': 'insufficient'}]}}
    # Minimal synthetic stored evidence, deliberately missing every metric.
    run = {'id': 12, 'params': {'service_model': 'recorded'}, 'result': {}}
    sections = shared_full_sections(des={'result': {'summary': {'served': 999}}}, mc=None, validation=None,
                                    decision=decision, counter_breaks=False, full_day=run)
    assert all(item['value'] is None for item in sections['simulation']['headlines'])
    checks = {item['check']: item for item in sections['checks']}
    assert 'service model recorded' in checks['Simulation run']['settings']
    assert 'seed not recorded' in checks['Simulation run']['settings']
    assert checks['Results check']['outcome'].startswith('0 passed, 0 failed, 1 without enough evidence')


@pytest.mark.parametrize('run', [None, {'id': 13}])
def test_a_full_day_report_never_falls_back_to_legacy_when_its_run_is_missing_or_wrong(run):
    with pytest.raises(ValueError, match='matching saved run'):
        shared_full_sections(des={'result': {'summary': {'served': 999}}}, mc=None, validation=None,
                             decision={'basis': 'full_day_run', 'evidence_ids': {'continuous': 12}},
                             counter_breaks=False, full_day=run)


def test_full_day_zero_values_are_preserved_without_substituting_a_different_wait_mean():
    run = {'id': 12, 'params': {}, 'result': {'summary': {
        'waiting_time': {'customer_weighted_mean_hours': {'value': 0},
                         'mean_of_replication_means_hours': {'mean': .5}},
        'metrics': {'max_queue_in_horizon': {'mean': 0}},
        'customers': {'per_replication': {'served': {'mean': 0}, 'unserved': {'mean': 0}}}},
        'period_utilization': {'segments': [{'segment_id': 'known-id', 'summary': {'mean': None}}]}}}
    before = deepcopy(run)
    sections = shared_full_sections(des=None, mc=None, validation=None, counter_breaks=False, full_day=run,
                                    decision={'basis': 'full_day_run', 'evidence_ids': {'continuous': 12}})
    assert [item['value'] for item in sections['simulation']['headlines']] == [0, 0, 0, 0]
    assert sections['simulation']['charts'][0]['points'] == [{'label': 'known-id', 'value': None}]
    assert sections['checks'][2]['outcome'] == 'not recorded by this Decision'
    assert run == before


def test_simulated_current_headers_fit_the_pdf_printable_page():
    headers = ['Period', 'Wait: Current staffing (simulated) (min)', 'Wait with the plan (min)',
               'Cost: Current staffing (simulated)', 'Cost with the plan']
    table = _table(_styles(), headers, [['08:00-09:00', '2.57', '1.59', '442.03', '442.03']])
    width, _ = table.wrap(letter[0] - 2 * inch, letter[1] - 2 * inch)
    assert width <= letter[0] - 2 * inch
